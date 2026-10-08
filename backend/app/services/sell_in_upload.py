"""
Reads the sell-in tracker workbook: one row per PO line, with the month of
sales, invoice year, customer, product description and quantity in packs. The
sheet is found by its columns, not its name. Pure -- no database.
inventory_service matches the lines to customers and SKUs and decides which
ones can be written.
"""

from datetime import date
from io import BytesIO
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

REQUIRED_COLUMNS = ("Month of Sales", "Invoice Year", "Customer Name", "DESCRIPTION", "QTY in Packs")
HEADER_SEARCH_ROWS = 10

MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "sept": 9, "oct": 10, "nov": 11, "dec": 12,
}

# The tracker's own spelling of each product, as the catalog names it.
DESCRIPTION_PRODUCTS = {
    "AIRE ADULT PANTS S/M 10S": "Aire Adult Pants S/M",
    "AIRE ADULT PANTS L 10S": "Aire Adult Pants L",
    "AIRE ADULT PANTS XL 10S": "Aire Adult Pants XL",
    "ADULT PULL UP PANTS ULTRA S/M": "Aire Ultra Pants S/M",
    "ADULT PULL UP PANTS ULTRA L": "Aire Ultra Pants L",
    "ADULT PULL UP PANTS ULTRA XL": "Aire Ultra Pants XL",
    "ADULT TAPE ULTRA S/M": "Aire Ultra Tape S/M",
    "ADULT TAPE ULTRA L": "Aire Ultra Tape L",
    "ADULT TAPE ULTRA XL": "Aire Ultra Tape XL",
}

_NO_ROWS_MESSAGE = (
    "The sell-in sheet has no rows to upload. Check that this is the sell-in tracker "
    "and that the sheet has data under its header row."
)


def product_for_description(description: str) -> str | None:
    return DESCRIPTION_PRODUCTS.get(description.upper())


def quantity_problem(qty: float | None) -> str | None:
    if qty is None:
        return "Quantity in packs is missing"
    if qty < 0:
        return "Quantity in packs is negative"
    return None


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _cell(row: tuple, index: int):
    return row[index] if index < len(row) else None


def _month_start(month_name: str, year) -> date | None:
    month = MONTHS.get(month_name.lower())
    if month is None or not isinstance(year, (int, float)):
        return None
    return date(int(year), month, 1)


def _find_columns(header: list[str]) -> dict[str, int] | None:
    """Column index for each required name, or None when the header lacks one."""

    columns = {}
    for name in REQUIRED_COLUMNS:
        for index, cell in enumerate(header):
            if name == "Invoice Year":
                matches = cell.lower().startswith(name.lower())
            else:
                matches = cell.lower() == name.lower()
            if matches:
                columns[name] = index
                break
        else:
            return None
    return columns


def _find_sheet(workbook) -> tuple:
    """
    The one worksheet whose first few rows hold the sell-in columns, as
    (worksheet, header row number, column indexes). Refused when none or
    several do, so the upload never reads the wrong sheet.
    """

    found = []
    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(max_row=HEADER_SEARCH_ROWS, values_only=True)
        for number, row in enumerate(rows, start=1):
            columns = _find_columns([_text(cell) for cell in row])
            if columns:
                found.append((worksheet, number, columns))
                break

    if not found:
        raise ValueError(
            "No sheet in this file has the sell-in columns (" + ", ".join(REQUIRED_COLUMNS) + ") "
            f"in its first {HEADER_SEARCH_ROWS} rows."
        )
    if len(found) > 1:
        names = ", ".join(f"'{worksheet.title}'" for worksheet, _, _ in found)
        raise ValueError(
            f"More than one sheet has the sell-in columns ({names}). Remove the extra sheet(s) and upload again."
        )
    return found[0]


def read_sell_in_lines(data: bytes) -> list[dict]:
    """
    One dict per non-empty row: customer_name, month (the first day of the month
    of sales in the invoice year, or None when either is unreadable),
    description and qty (quantity in packs, or None when it is not a number).
    """

    try:
        workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException):
        raise ValueError("Invalid file. Please upload the correct file.")
    worksheet, header_number, columns = _find_sheet(workbook)

    lines = []
    for row in worksheet.iter_rows(min_row=header_number + 1, values_only=True):
        customer_name = _text(_cell(row, columns["Customer Name"]))
        month_name = _text(_cell(row, columns["Month of Sales"]))
        if not customer_name and not month_name:
            continue
        qty = _cell(row, columns["QTY in Packs"])
        lines.append(
            {
                "customer_name": customer_name,
                "month": _month_start(month_name, _cell(row, columns["Invoice Year"])),
                "description": _text(_cell(row, columns["DESCRIPTION"])),
                "qty": float(qty) if isinstance(qty, (int, float)) else None,
            }
        )
    if not lines:
        raise ValueError(_NO_ROWS_MESSAGE)
    return lines
