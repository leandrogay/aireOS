"""
Find the data table inside an uploaded retailer file.

Retailer files are rarely a clean rectangle starting at A1. NHG's workbook
opens on a pivot table; Sheng Siong's claim statement has a title, a block of
"Key : value" lines, a header on row 7, merged date cells, a blank spacer
column and footer totals. Everything downstream -- the mapping proposal, the
checks and the load -- needs the same answer to "where is the data?", so it is
worked out once, here, and returned with the table.

Pure apart from parsing the bytes: no GCS, no database, no model.
"""

import csv
import datetime
import io
import re
from pathlib import Path

import openpyxl
import pandas as pd


class LayoutError(ValueError):
    """No usable data table could be found in the file."""


# A header row is looked for in the first rows only; a sheet whose header sits
# lower than this is not a layout we have seen.
HEADER_SCAN_ROWS = 30

# Labels Excel writes into a pivot table. A sheet carrying them is a summary of
# the data, never the data itself. "Grand Total" alone is not one: an ordinary
# data sheet can end with a grand-total footer.
PIVOT_MARKERS = ("row labels", "column labels")
PIVOT_MEASURE = re.compile(r"^(sum|count|average) of (?P<column>.+)$", re.IGNORECASE)

# "Total", "Grand Total", "Total Claim before GST" -- a row whose label starts
# like this, at the bottom of the table, is a footer.
FOOTER_LABEL = re.compile(r"^(grand\s+)?total\b", re.IGNORECASE)
# Exactly "Total" / "Subtotal": a row inside the table labelled this way is a
# subtotal and would double count if it stayed.
SUBTOTAL_LABEL = re.compile(r"^(grand\s+|sub\s*)?total$", re.IGNORECASE)
TOTAL_COLUMN = re.compile(r"^(grand\s+)?total$", re.IGNORECASE)

# "CLAIM SALES SUMMARY FOR 1/6/2026 TO 30/6/2026" (Sheng Siong) or
# "DATE : 01 Aug 2026 TO 31 Aug 2026" (SingHealth).
_STATED_DATE = r"\d{1,2}[/-]\d{1,2}[/-]\d{4}|\d{1,2}[ -][A-Za-z]{3,9}[ -]\d{4}"
STATED_PERIOD = re.compile(
    rf"({_STATED_DATE})\s*(?:to|-|–)\s*({_STATED_DATE})",
    re.IGNORECASE,
)
NUMBER_LIKE = re.compile(r"^-?\d+(\.\d+)?$")
DATE_LIKE = re.compile(r"^\d{4}-\d{2}-\d{2}")

EXCEL_EXTENSIONS = (".xlsx", ".xlsm")
TEXT_EXTENSIONS = (".csv", ".txt")


# ---- Reading cells -----------------------------------------------------------

def _cell_text(value) -> str | None:
    """
    One cell as text, the way the rest of the pipeline expects it.

    Everything is text so identifiers keep their exact form (a store code of
    "0420" must not become 420). Dates become ISO strings and whole numbers
    lose the ".0" Excel gives them; an empty or blank cell is None.
    """
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value).upper()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.15g}"

    text = str(value).strip()
    return text or None


def _excel_sheets(data: bytes) -> list[dict]:
    """
    Every worksheet as a grid of text cells, merged ranges already filled down.

    A merged cell holds its value only in its top-left corner; Sheng Siong's
    BusinessDate spans each date's block of rows, so every row but the first
    would read as dated nowhere. Copying the corner value down the range is
    what the merge meant. Only down: a title merged across a row is still one
    cell, and copying it sideways would make it look like a header row.
    """
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise LayoutError(f"Could not open the workbook: {exc}") from exc

    sheets = []
    for worksheet in workbook.worksheets:
        rows = [
            [_cell_text(value) for value in row]
            for row in worksheet.iter_rows(values_only=True)
        ]
        filled = set()
        for merged in worksheet.merged_cells.ranges:
            top, left = merged.min_row - 1, merged.min_col - 1
            if merged.max_row == merged.min_row:
                continue
            if top >= len(rows) or left >= len(rows[top]):
                continue
            anchor = rows[top][left]
            for row_index in range(top + 1, min(merged.max_row, len(rows))):
                if left < len(rows[row_index]):
                    rows[row_index][left] = anchor
            filled.add((top, left))

        sheets.append({
            "name": worksheet.title,
            "rows": rows,
            # openpyxl keeps a sheet's pivot tables here; private, so the
            # content markers below are the check that does not depend on it.
            "has_pivot": bool(getattr(worksheet, "_pivots", None)),
            "merged": filled,
        })
    return sheets


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _text_sheet(filename: str, data: bytes) -> dict:
    """
    A delimited file as a one-sheet grid.

    Read with the csv module rather than pandas so a preamble line with fewer
    fields than the header does not abort the read. A .txt file's delimiter is
    sniffed (FairPrice exports are tab separated); a .csv is comma separated.
    """
    text = _decode(data)
    if Path(filename).suffix.lower() == ".txt":
        try:
            dialect = csv.Sniffer().sniff(text[:20000], delimiters="\t,;|")
        except csv.Error:
            dialect = csv.excel_tab
    else:
        dialect = csv.excel

    rows = [
        [_cell_text(value) for value in row]
        for row in csv.reader(io.StringIO(text), dialect)
    ]
    return {"name": None, "rows": rows, "has_pivot": False, "merged": set()}


# ---- Finding the table on a sheet --------------------------------------------

def _filled(row: list) -> list[str]:
    return [value for value in row if value is not None]


def _is_number_or_date(value: str) -> bool:
    return bool(NUMBER_LIKE.match(value) or DATE_LIKE.match(value))


def _find_header_row(rows: list[list]) -> int | None:
    """
    The first row that looks like a header: at least half as wide as the
    widest row, mostly words rather than numbers or dates, and followed by at
    least one row of data. Titles and "Key : value" lines are too narrow;
    data rows are mostly numbers.
    """
    scanned = rows[:HEADER_SCAN_ROWS]
    widths = [len(_filled(row)) for row in rows[: HEADER_SCAN_ROWS * 4]]
    widest = max(widths, default=0)
    if widest == 0:
        return None

    for index, row in enumerate(scanned):
        cells = _filled(row)
        if len(cells) < max(2, widest / 2):
            continue
        words = [value for value in cells if not _is_number_or_date(value)]
        if len(words) < 0.8 * len(cells):
            continue
        if any(_filled(later) for later in rows[index + 1 :]):
            return index
    return None


def _clean_headers(row: list, width: int) -> list[str]:
    """
    Header text with whitespace normalised, blanks named and repeats numbered.

    Sheng Siong writes "BusinessDate " and "Total " with trailing spaces; a
    contract naming "Total" must still find it. A blank header gets pandas'
    "Unnamed: n" so every column has a name.
    """
    headers = []
    seen: dict[str, int] = {}
    for index in range(width):
        value = row[index] if index < len(row) else None
        name = " ".join(str(value).split()) if value is not None else ""
        if not name:
            name = f"Unnamed: {index}"
        if name in seen:
            seen[name] += 1
            name = f"{name}.{seen[name]}"
        else:
            seen[name] = 0
        headers.append(name)
    return headers


def _parse_preamble(rows: list[list]) -> dict:
    """
    The title and "Key : value" lines above the header, and any period they
    state. Kept for reconciliation and for the reviewer; never mapped.
    """
    title = None
    metadata: dict[str, str] = {}

    for row in rows:
        cells = _filled(row)
        if not cells:
            continue
        if len(cells) == 1:
            text = cells[0]
            if ":" in text and not title:
                key, _, value = text.partition(":")
                if key.strip() and value.strip():
                    metadata[key.strip()] = value.strip()
                    continue
            title = title or text
        elif len(cells) == 2:
            key, value = cells
            metadata[key.strip().rstrip(":").strip()] = value.lstrip(": ").strip()

    stated_period = None
    for text in [title or "", *metadata.values()]:
        match = STATED_PERIOD.search(text)
        if match:
            start = pd.to_datetime(match.group(1), dayfirst=True, errors="coerce")
            end = pd.to_datetime(match.group(2), dayfirst=True, errors="coerce")
            if pd.notna(start) and pd.notna(end):
                stated_period = {
                    "start": start.date().isoformat(),
                    "end": end.date().isoformat(),
                    "stated_in": text,
                }
                break

    return {"title": title, "metadata": metadata, "stated_period": stated_period}


def _rows_above_header(rows: list[list], headers: list[str], keep: list[int]) -> list[dict]:
    """
    The cells above the header, keyed by the header of the column they sit
    in. SingHealth writes each outlet's code in the row above its name, so the
    header alone does not say everything about a column.
    """
    above = []
    for index, row in enumerate(rows):
        cells = {
            headers[column]: row[column]
            for column in keep
            if column < len(row) and row[column] is not None
        }
        if cells:
            above.append({"row": index + 1, "cells": cells})
    return above


def _footer_label(row: list) -> str | None:
    """The "Total ..." label that marks a footer row, if the row has one."""
    return next((value for value in row if value and FOOTER_LABEL.match(value)), None)


def _summary_label(row: list) -> str | None:
    """
    The label of a summary line: text that starts the row past its first
    column, the way "Supplier Claim before GST" sits under Sheng Siong's
    outlet columns. A data row starts in the first column.
    """
    if row[0] is not None:
        return None
    first = next((value for value in row if value is not None), None)
    if first is None or _is_number_or_date(first):
        return None
    return first


def _pivot_reason(sheet: dict, header_row: int | None) -> str | None:
    """
    Why a sheet is a pivot table, or None when it is not one: Excel's own pivot
    object, a "Row Labels" cell, or a "Sum of …" measure heading the table.
    """
    if sheet["has_pivot"]:
        return "pivot table"
    rows = sheet["rows"]
    for row in rows[: HEADER_SCAN_ROWS * 2]:
        if any(value.strip().lower() in PIVOT_MARKERS for value in _filled(row)):
            return "pivot table"
    if header_row is not None:
        if any(PIVOT_MEASURE.match(value.strip()) for value in _filled(rows[header_row - 1])):
            return "pivot table"
    return None


def _pivot_totals(sheet: dict) -> list[dict]:
    """
    A pivot sheet's grand totals, keyed by the column it summarises.

    NHG's first sheet sums "Qty in Un. of Entry" to -1,868; the data sheet's
    output should net to the same figure, so the total is kept to check against.
    """
    rows = sheet["rows"]
    measures: list[tuple[int, str]] = []
    totals = []
    for row in rows:
        for index, value in enumerate(row):
            if value is None:
                continue
            match = PIVOT_MEASURE.match(value.strip())
            if match:
                measures.append((index, match.group("column").strip()))
        cells = _filled(row)
        if cells and cells[0].strip().lower() == "grand total":
            for index, column in measures:
                if index < len(row) and row[index] is not None:
                    try:
                        totals.append({
                            "sheet": sheet["name"],
                            "column": column,
                            "total": float(row[index]),
                        })
                    except ValueError:
                        pass
    return totals


def _table_on_sheet(sheet: dict) -> dict | None:
    """Locate the header, data rows, footers and special columns on one sheet."""
    rows = sheet["rows"]
    header_index = _find_header_row(rows)
    if header_index is None:
        return None

    body = rows[header_index + 1 :]
    width = max(
        [len(rows[header_index])] + [len(row) for row in body],
        default=0,
    )
    headers = _clean_headers(rows[header_index], width)

    def padded(row):
        return (list(row) + [None] * width)[:width]

    indexed = [(header_index + 1 + offset, padded(row)) for offset, row in enumerate(body)]
    indexed = [(number, row) for number, row in indexed if _filled(row)]

    # Footers: the rows at the very bottom from the last "Total" row down.
    # Sheng Siong writes its claim lines under the total -- "Total Claim before
    # GST" in older statements, "Supplier Claim before GST" in newer ones -- so
    # a summary line below a total is a footer whatever it is labelled.
    # Without a total above them, summary lines stay data.
    tail = len(indexed)
    while tail and _summary_label(indexed[tail - 1][1]) and not _footer_label(indexed[tail - 1][1]):
        tail -= 1
    start = tail
    while start and _footer_label(indexed[start - 1][1]):
        start -= 1
    if start == tail:
        start = len(indexed)
    footer_rows = [
        {
            "label": _footer_label(row) or _summary_label(row),
            "values": {headers[i]: value for i, value in enumerate(row) if value is not None},
        }
        for _, row in indexed[start:]
    ]
    indexed = indexed[:start]

    subtotal_rows = 0
    data_rows = []
    for _, row in indexed:
        if any(value and SUBTOTAL_LABEL.match(value) for value in row):
            subtotal_rows += 1
            continue
        data_rows.append(row)

    columns_with_data = {
        index for row in data_rows for index, value in enumerate(row) if value is not None
    }
    keep = [
        index
        for index, header in enumerate(headers)
        if index in columns_with_data or not header.startswith("Unnamed: ")
    ]
    kept_headers = [headers[index] for index in keep]

    spacer_columns = [headers[index] for index in keep if index not in columns_with_data]
    total_columns = [header for header in kept_headers if TOTAL_COLUMN.match(header)]
    filled_columns = sorted({
        headers[left]
        for top, left in sheet["merged"]
        if top > header_index and left < width
    })

    frame = pd.DataFrame(
        [[row[index] for index in keep] for row in data_rows],
        columns=kept_headers,
        dtype=object,
    )
    frame = frame.astype(object).where(frame.notna(), None)

    return {
        "frame": frame,
        "header_row": header_index + 1,
        "data_rows": len(data_rows),
        "footer_rows": footer_rows,
        "subtotal_rows": subtotal_rows,
        "total_columns": total_columns,
        "spacer_columns": spacer_columns,
        "filled_columns": filled_columns,
        "rows_above_header": _rows_above_header(rows[:header_index], headers, keep),
        **_parse_preamble(rows[:header_index]),
    }


# ---- The one function callers use ---------------------------------------------

def read_sheets(filename: str, data: bytes) -> list[dict]:
    """Every sheet of an upload as a grid of text cells."""
    extension = Path(filename).suffix.lower()
    if extension in EXCEL_EXTENSIONS:
        return _excel_sheets(data)
    if extension in TEXT_EXTENSIONS:
        return [_text_sheet(filename, data)]
    raise LayoutError(f"Unsupported file format for mapping: {extension or '(none)'}")


def read_table(filename: str, data: bytes) -> tuple[pd.DataFrame, dict]:
    """
    The file's data table as text cells, plus a description of where it was.

    Of a workbook's sheets the one chosen is the non-pivot sheet with the most
    data rows; every other sheet is listed in `rejected_sheets` with the
    reason. The layout also carries what was set aside so it can be checked
    against: the preamble and any period it states, the cells above the
    header by column, footer rows, total and spacer columns, and a pivot
    sheet's grand totals.
    """
    sheets = read_sheets(filename, data)

    chosen = None
    rejected = []
    pivot_totals = []
    for sheet in sheets:
        table = _table_on_sheet(sheet)
        reason = _pivot_reason(sheet, table and table["header_row"])
        if reason:
            rejected.append({"name": sheet["name"], "reason": reason})
            pivot_totals.extend(_pivot_totals(sheet))
            continue
        if table is None or not table["data_rows"]:
            rejected.append({"name": sheet["name"], "reason": "no data table found"})
            continue
        if chosen is None or table["data_rows"] > chosen[1]["data_rows"]:
            if chosen is not None:
                rejected.append({"name": chosen[0]["name"], "reason": "fewer data rows"})
            chosen = (sheet, table)
        else:
            rejected.append({"name": sheet["name"], "reason": "fewer data rows"})

    if chosen is None:
        reasons = ", ".join(
            f"{item['name'] or 'the file'}: {item['reason']}" for item in rejected
        )
        raise LayoutError(f"No data table found ({reasons or 'the file is empty'}).")

    sheet, table = chosen
    frame = table.pop("frame")
    layout = {
        "sheet": sheet["name"],
        "sheets": [item["name"] for item in sheets],
        "rejected_sheets": rejected,
        "pivot_totals": pivot_totals,
        **table,
    }
    return frame, layout
