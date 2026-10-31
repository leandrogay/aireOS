"""Read retailer calendars into a reviewable JSON draft; never write SQL."""

import argparse
import csv
import json
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.schemas.catalog import normalise_name  # noqa: E402
from app.services.promotion_pricing import mechanic_multiplier, voucher_multiplier  # noqa: E402


_SKIP_SHEETS = {
    "Instructions", "Settings", "All Retailers", "Display Tracker", "Promoter Tracker",
}
_SKU_RANGE_PROPOSALS = {
    "Flagship": "Aire Adult Diaper Pants",
    "Ultra Pants": "Aire Adult Diaper Ultra Pants",
    "Ultra Tape": "Aire Adult Diaper Ultra Tape",
}
_RETAILER_PROPOSALS = {
    # Match the existing retailer records; the customer slug is different.
    "Sheng Siong Offline": "sheng_siong_offline",
    "Sheng Siong Online": "sheng_siong_online",
    "NHG": "nhg_offline",
    "Prime": "prime_offline",
    # The unqualified SingHealth calendar covers both customer channels.
    "SingHealth": "singhealth",
}
_MECHANIC_COLUMNS = {
    "promotion_mechanic": "monthly",
    "promotion_mechanic_2": "monthly",
    "side_offer_mechanic": "side_offer",
    "side_offer_mechanic_2": "side_offer",
}


def _calendar_date(value) -> str | None:
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return None


def preview_calendar(path: Path) -> list[dict]:
    """Extract filled mechanics, keeping the exact sheet/row for review."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    candidates = []
    try:
        for sheet in workbook:
            if sheet.title in _SKIP_SHEETS or sheet.title.casefold().startswith("fairprice"):
                continue
            headers = None
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                if not values:
                    continue
                if values[0] in ("period_start", "week_start"):
                    headers = [str(value).strip() if value is not None else "" for value in values]
                    continue
                if not headers:
                    continue
                data = dict(zip(headers, values))
                start = _calendar_date(data.get("period_start") or data.get("week_start"))
                end = _calendar_date(data.get("period_end") or data.get("week_end"))
                if not start or not end:
                    continue
                source_range = str(data.get("sku_range") or "").strip()
                voucher = str(data.get("voucher") or "").strip() or None
                filled = [
                    (column, promo_type, str(data.get(column) or "").strip())
                    for column, promo_type in _MECHANIC_COLUMNS.items()
                    if str(data.get(column) or "").strip()
                ]
                if data.get("promotion_mechanic_2"):
                    filled = [item for item in filled if item[0] != "promotion_mechanic"]
                if data.get("side_offer_mechanic_2"):
                    filled = [item for item in filled if item[0] != "side_offer_mechanic"]
                pack_columns = ("promotion_mechanic", "promotion_mechanic_2",
                                "side_offer_mechanic", "side_offer_mechanic_2")
                if voucher and not any(column in pack_columns for column, _, _ in filled):
                    promo_type = "monthly" if data.get("period_start") else "side_offer"
                    filled.append(("voucher_only", promo_type, "Voucher"))
                for column, promo_type, mechanic in filled:
                    multiplier = mechanic_multiplier(mechanic)
                    voucher_rate = voucher_multiplier(voucher)
                    issues = []
                    if multiplier is None:
                        issues.append("Mechanic needs manual interpretation")
                    if voucher_rate is None:
                        issues.append("Voucher needs manual interpretation")
                    if source_range not in _SKU_RANGE_PROPOSALS:
                        issues.append("SKU range has no proposed catalog match")
                    candidates.append({
                        "approved": False,
                        "source": path.name,
                        "sheet": sheet.title,
                        "row": row_number,
                        "mechanic_column": column,
                        "retailer_proposal": _RETAILER_PROPOSALS.get(
                            sheet.title, normalise_name(sheet.title)
                        ),
                        "period_start": start,
                        "period_end": end,
                        "period_label": str(data.get("period_label") or data.get("week_label") or "").strip() or None,
                        "promo_type": promo_type,
                        "source_sku_range": source_range,
                        "sku_range_proposal": _SKU_RANGE_PROPOSALS.get(source_range),
                        "promotion_mechanic": mechanic,
                        "voucher": voucher,
                        "estimated_multiplier": (
                            str(multiplier * voucher_rate)
                            if multiplier is not None and voucher_rate is not None else None
                        ),
                        "store_scope": "all existing stores under retailer; verify against catalog before import",
                        "review_note": (
                            "Second mechanic replaces the first on this row"
                            if column.endswith("_2") else None
                        ),
                        "replaced_mechanic": (
                            str(data.get(column.removesuffix("_2")) or "").strip() or None
                            if column.endswith("_2") else None
                        ),
                        "issues": issues,
                    })
    finally:
        workbook.close()
    return candidates


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbooks", nargs="+", type=Path)
    parser.add_argument("--output", type=Path, help="Write the full JSON review draft here")
    args = parser.parse_args()

    candidates = [row for path in args.workbooks for row in preview_calendar(path)]
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(candidates, indent=2), encoding="utf-8")
        csv_path = args.output.with_suffix(".csv")
        columns = (
            "approved", "source", "sheet", "row", "mechanic_column",
            "retailer_proposal", "period_start", "period_end", "period_label",
            "promo_type", "source_sku_range", "sku_range_proposal",
            "promotion_mechanic", "replaced_mechanic", "voucher",
            "estimated_multiplier", "review_note", "issues",
        )
        with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
            writer.writeheader()
            for row in candidates:
                writer.writerow({**row, "issues": "; ".join(row["issues"])})

    by_sheet = Counter((row["source"], row["sheet"]) for row in candidates)
    issues = Counter(issue for row in candidates for issue in row["issues"])
    print(f"{len(candidates)} filled promotion mechanics; no SQL writes")
    for (source, sheet), count in sorted(by_sheet.items()):
        print(f"  {source} / {sheet}: {count}")
    for issue, count in issues.items():
        print(f"  review: {issue}: {count}")
    if args.output:
        print(f"Review drafts: {args.output}, {csv_path}")


if __name__ == "__main__":
    main()
