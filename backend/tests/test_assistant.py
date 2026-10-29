import io
import json
import base64
import datetime
from types import SimpleNamespace

import pandas as pd
import pytest
from conftest import FrozenDatetime

from google.api_core.exceptions import ServiceUnavailable

from app.services import assistant, forecast_summary, inventory_service, promotion_service, sellout_service


def _freeze_clock(monkeypatch):
    monkeypatch.setattr(
        assistant, "datetime", SimpleNamespace(datetime=FrozenDatetime, date=datetime.date,
                                               timedelta=datetime.timedelta, timezone=datetime.timezone),
    )


# ---- Fakes mirroring just the SDK response shape assistant.py reads --------

class FakeFunctionCall:
    def __init__(self, id, name, args):
        self.id = id
        self.name = name
        self.args = args


class FakeCandidate:
    def __init__(self, content, finish_reason="STOP"):
        self.content = content
        self.finish_reason = finish_reason


class FakeResponse:
    """content is a plain dict placeholder for candidate.content -- ask()
    only appends it to `contents` for history, it never parses it back out
    (function calls come from the separate `function_calls` property, text
    from the separate `text` property, matching the real SDK)."""

    def __init__(self, function_calls=None, text=None, finish_reason="STOP", content=None):
        self.function_calls = function_calls or []
        self.text = text
        self.candidates = [FakeCandidate(content or {"role": "model", "parts": []}, finish_reason)]


class FakeModels:
    """Returns each queued response in order, one per call to .generate_content()."""

    def __init__(self, responses):
        self._responses = list(responses)
        self.calls = []

    def generate_content(self, **kwargs):
        # Snapshot contents -- assistant.ask() mutates the same list object
        # across loop iterations, so without a copy every recorded call
        # would end up pointing at the final, fully-mutated list.
        snapshot = {**kwargs, "contents": list(kwargs["contents"])}
        self.calls.append(snapshot)
        return self._responses.pop(0)


class FakeClient:
    def __init__(self, responses):
        self.models = FakeModels(responses)


def _install_fake_client(monkeypatch, responses):
    fake_client = FakeClient(responses)
    monkeypatch.setattr(assistant, "get_client", lambda: fake_client)
    return fake_client


def _final_answer(
    text="Revenue was $1,000 last week.",
    grounded=True,
    data_source="aire_data",
    follow_ups=None,
    export_format="none",
    chart_style=None,
    uses_previous_chart=False,
):
    payload = {
        "answer": text,
        "grounded": grounded,
        "data_source": data_source,
        "follow_up_prompts": follow_ups or ["Break this down by store?", "Compare to last month?"],
        "uses_previous_chart": uses_previous_chart,
        "export_format": export_format,
        "chart_style": chart_style or dict(assistant._DEFAULT_CHART_STYLE),
    }
    return FakeResponse(
        text=json.dumps(payload),
        content={"role": "model", "parts": [{"text": json.dumps(payload)}]},
    )


def _tool_call_response(*calls):
    """calls: (id, name, args) tuples -- one FakeFunctionCall per call."""
    function_calls = [FakeFunctionCall(id, name, args) for id, name, args in calls]
    placeholder_parts = [{"function_call": {"id": id, "name": name, "args": args}} for id, name, args in calls]
    return FakeResponse(function_calls=function_calls, content={"role": "model", "parts": placeholder_parts})


def _dump(content):
    """Same logic as assistant._serialize_contents for one item -- real
    types.Content objects (built by ask() itself for tool-result turns) need
    .model_dump() to inspect; plain dicts (from our fakes) pass through."""
    return content.model_dump(exclude_none=True, mode="json") if hasattr(content, "model_dump") else content


# ---- ask(): tool loop shape -------------------------------------------------

def test_end_turn_immediately_with_no_tool_calls(monkeypatch):
    # grounded=False from the business-tools pass triggers the search
    # fallback (see the dedicated search-fallback tests below) -- here that
    # also comes back ungrounded, so the original answer wins.
    _install_fake_client(
        monkeypatch,
        [
            _final_answer(grounded=False, data_source="none", text="I don't have data for that."),
            _search_answer("Nothing relevant turned up."),
            _final_answer(grounded=False, data_source="none", text="Search didn't find anything either."),
        ],
    )

    result = assistant.ask("What's the meaning of life?", history=None, customer="fairprice")

    assert result["grounded"] is False
    assert result["data_source"] == "none"
    assert result["has_chart"] is False
    assert result["has_table"] is False


def test_first_request_sends_question_as_user_message(monkeypatch):
    fake_client = _install_fake_client(monkeypatch, [_final_answer()])

    assistant.ask("How did we do last week?", history=None, customer="fairprice")

    first_call = fake_client.models.calls[0]
    assert first_call["contents"][-1] == {"role": "user", "parts": [{"text": "How did we do last week?"}]}


def test_prior_history_is_carried_forward(monkeypatch):
    fake_client = _install_fake_client(monkeypatch, [_final_answer()])
    history = [
        {"role": "user", "parts": [{"text": "hi"}]},
        {"role": "model", "parts": [{"text": "hello"}]},
    ]

    assistant.ask("follow up question", history=history, customer="fairprice")

    first_call = fake_client.models.calls[0]
    assert first_call["contents"][0] == history[0]
    assert first_call["contents"][1] == history[1]


def test_returned_messages_can_be_replayed_next_turn(monkeypatch):
    _install_fake_client(monkeypatch, [_final_answer()])

    result = assistant.ask("How did we do last week?", history=None, customer="fairprice")

    # messages ends with the model's final turn appended (Gemini's role for
    # the model's own turns is "model", not Anthropic's "assistant")
    assert result["messages"][-1]["role"] == "model"


# ---- ask(): function_call -> function_response round trip ------------------

def test_tool_call_result_is_fed_back_and_chart_is_built(monkeypatch):
    summary = {
        "offline": {
            "storeFormats": [],
            "periodTotal": [
                {"period_label": "Week 1", "period_start": "2026-08-01", "revenue": 1000.0, "units": 10},
                {"period_label": "Week 2", "period_start": "2026-08-08", "revenue": 1500.0, "units": 15},
            ],
            "periodByFormat": [],
        },
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "week"}))
    fake_client = _install_fake_client(monkeypatch, [tool_call, _final_answer()])

    result = assistant.ask("How is revenue trending?", history=None, customer="fairprice")

    assert result["has_chart"] is True
    assert result["chart_categories"] == ["Week 1", "Week 2"]
    assert result["chart_series"] == [{"label": "Revenue", "values": [1000.0, 1500.0]}]

    # the second generate_content() call carries the function result back
    second_call_contents = fake_client.models.calls[1]["contents"]
    tool_result_message = _dump(second_call_contents[-1])
    assert tool_result_message["role"] == "user"
    function_response = tool_result_message["parts"][0]["function_response"]
    assert function_response["name"] == "get_sales_summary"
    assert "result" in function_response["response"]


def test_get_sales_summary_chart_combines_offline_and_online(monkeypatch):
    # Regression: when the model answers from get_sales_summary alone
    # (skipping compare_periods), a headline like "$25,915 vs $23,703" is
    # the ALL-CHANNEL total -- the chart must sum offline+online per
    # period, not show just one channel.
    summary = {
        "offline": {
            "storeFormats": [],
            "periodTotal": [
                {"period_label": "May 2026", "period_start": "2026-05-01", "revenue": 15430.44, "units": 1561},
                {"period_label": "June 2026", "period_start": "2026-06-01", "revenue": 16601.92, "units": 1735},
            ],
            "periodByFormat": [],
        },
        "online": {
            "storeFormats": [],
            "periodTotal": [
                {"period_label": "May 2026", "period_start": "2026-05-01", "revenue": 8272.36, "units": 880},
                {"period_label": "June 2026", "period_start": "2026-06-01", "revenue": 9313.36, "units": 1007},
            ],
            "periodByFormat": [],
        },
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="June was up on May: $25,915.28 vs $23,702.80.")])

    result = assistant.ask("compare May to June", history=None, customer="fairprice")

    assert result["chart_categories"] == ["May 2026", "June 2026"]
    assert result["chart_series"] == [{"label": "Revenue", "values": [23702.80, 25915.28]}]


def test_tool_error_is_reported_as_is_error_and_loop_continues(monkeypatch):
    def _boom(**kwargs):
        raise ValueError("granularity must be one of ('week', 'month')")

    monkeypatch.setattr(sellout_service, "get_dashboard_summary", _boom)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "decade"}))
    fake_client = _install_fake_client(
        monkeypatch,
        [
            tool_call,
            _final_answer(grounded=False, data_source="none"),
            _search_answer("Nothing relevant turned up."),
            _final_answer(grounded=False, data_source="none"),
        ],
    )

    result = assistant.ask("bad question", history=None, customer="fairprice")

    second_call_contents = fake_client.models.calls[1]["contents"]
    function_response = _dump(second_call_contents[-1])["parts"][0]["function_response"]["response"]
    assert "error" in function_response
    assert "Invalid arguments" in function_response["error"]
    # a failed tool call never becomes a chart
    assert result["has_chart"] is False


def test_multiple_tool_use_blocks_in_one_turn_all_get_results(monkeypatch):
    monkeypatch.setattr(sellout_service, "get_customer_options", lambda: [{"value": "fairprice", "label": "Fairprice"}])
    monkeypatch.setattr(sellout_service, "get_store_options", lambda customer: [{"store_code": "S1", "store_name": "Store 1"}])

    tool_call = _tool_call_response(
        ("tu_1", "list_customers", {}),
        ("tu_2", "list_stores", {"customer": "fairprice"}),
    )
    fake_client = _install_fake_client(monkeypatch, [tool_call, _final_answer()])

    assistant.ask("what stores do we have", history=None, customer="fairprice")

    tool_results = _dump(fake_client.models.calls[1]["contents"][-1])["parts"]
    names = {p["function_response"]["name"] for p in tool_results}
    assert names == {"list_customers", "list_stores"}


def test_get_sales_summary_mode_scoped_call_does_not_combine_channels(monkeypatch):
    # Regression (found via live testing): asked "offline sales in August",
    # the model correctly called get_sales_summary(mode="offline") and cited
    # the offline-only total ($5,473.15) -- but get_dashboard_summary always
    # computes BOTH channels regardless of what was asked, so the chart
    # (built by summing whatever channels are present in the raw result)
    # showed offline+online combined ($9,632.45), silently contradicting
    # the answer's own channel-scoped figure. _run_business_tool must strip
    # the non-requested channel before it ever reaches chart-building.
    full_summary = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-06", "revenue": 5473.15, "units": 388}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-06", "revenue": 4159.30, "units": 300}], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: full_summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month", "mode": "offline"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="Fairprice's offline sales in August 2026 were $5,473.15.")])

    result = assistant.ask("what was sales for fairprice offline in august 2026", history=None, customer="fairprice")

    assert result["chart_categories"] == ["August 2026"]
    assert result["chart_series"] == [{"label": "Revenue", "values": [5473.15]}]


def test_get_sales_summary_no_mode_still_combines_channels(monkeypatch):
    # The combining behaviour itself is correct and still needed when no
    # mode was specified (e.g. "total revenue across both channels").
    full_summary = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-06", "revenue": 5473.15, "units": 388}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-06", "revenue": 4159.30, "units": 300}], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: full_summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="Total sales in August 2026 were $9,632.45.")])

    result = assistant.ask("what was total sales in august 2026", history=None, customer="fairprice")

    assert result["chart_series"] == [{"label": "Revenue", "values": [9632.45]}]


def test_default_customer_used_when_tool_omits_it(monkeypatch):
    captured = {}

    def _fake_options(customer):
        captured["customer"] = customer
        return []

    monkeypatch.setattr(sellout_service, "get_store_options", _fake_options)

    tool_call = _tool_call_response(("tu_1", "list_stores", {}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer()])

    assistant.ask("stores?", history=None, customer="fairprice")

    assert captured["customer"] == "fairprice"


# ---- ask(): compare_periods / rank_skus chart shaping -----------------------

def test_compare_periods_builds_two_bar_chart(monkeypatch):
    comparison = {
        "current": {"start": "2026-08-17", "end": "2026-08-23", "revenue": 1200.0, "units": 60},
        "previous": {"start": "2026-08-10", "end": "2026-08-16", "revenue": 900.0, "units": 45, "available": True},
    }
    monkeypatch.setattr(sellout_service, "get_period_comparison", lambda **kwargs: comparison)

    tool_call = _tool_call_response(("tu_1", "compare_periods", {"comparison_type": "wow"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer()])

    result = assistant.ask("how does this week compare", history=None, customer="fairprice")

    assert result["chart_categories"] == ["Previous", "Current"]
    assert result["chart_series"] == [{"label": "Revenue", "values": [900.0, 1200.0]}]


def test_compare_periods_mode_scoped_calls_are_summed_to_a_combined_total(monkeypatch):
    # Regression: the model sometimes calls compare_periods once per
    # channel (mode="offline", mode="online") instead of once unscoped, to
    # narrate a per-channel breakdown -- with no unscoped call, whichever
    # scoped call happened to run last would show only that channel's
    # numbers under a combined-total headline. The two scoped calls must
    # be summed back into the same combined total instead.
    results = {
        "online": {
            "current": {"start": "2026-06-01", "end": "2026-06-30", "revenue": 9313.36, "units": 1007},
            "previous": {"start": "2026-05-01", "end": "2026-05-31", "revenue": 8272.36, "units": 880, "available": True},
        },
        "offline": {
            "current": {"start": "2026-06-01", "end": "2026-06-30", "revenue": 16601.92, "units": 1735},
            "previous": {"start": "2026-05-01", "end": "2026-05-31", "revenue": 15430.44, "units": 1561, "available": True},
        },
    }
    monkeypatch.setattr(sellout_service, "get_period_comparison", lambda **kwargs: results[kwargs["mode"]])

    # Confirmed live: the real model passes explicit previous_start/
    # previous_end for a genuine comparison (not just current_start/end) --
    # that's also the signal _build_chart uses to tell "real comparison"
    # apart from "compare_periods used as a plain range-total calculator".
    tool_call = _tool_call_response(
        ("tu_1", "compare_periods", {"current_start": "2026-06-01", "current_end": "2026-06-30", "previous_start": "2026-05-01", "previous_end": "2026-05-31", "mode": "online"}),
        ("tu_2", "compare_periods", {"current_start": "2026-06-01", "current_end": "2026-06-30", "previous_start": "2026-05-01", "previous_end": "2026-05-31", "mode": "offline"}),
    )
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="June was up on May: $25,915.28 vs $23,702.80.")])

    result = assistant.ask("compare May to June by channel", history=None, customer="fairprice")

    assert result["chart_categories"] == ["Previous", "Current"]
    assert result["chart_series"] == [{"label": "Revenue", "values": [23702.80, 25915.28]}]


def test_compare_periods_prefers_unscoped_call_over_mode_scoped_ones(monkeypatch):
    # Whenever an unscoped (combined) call exists, it's already the
    # all-channel total -- use it directly rather than trying to sum the
    # scoped calls too (which, called alongside it, would double-count).
    unscoped = {
        "current": {"start": "2026-06-01", "end": "2026-06-30", "revenue": 25915.28, "units": 2742},
        "previous": {"start": "2026-05-01", "end": "2026-05-31", "revenue": 23702.80, "units": 2441, "available": True},
    }
    offline_only = {
        "current": {"start": "2026-06-01", "end": "2026-06-30", "revenue": 16601.92, "units": 1735},
        "previous": {"start": "2026-05-01", "end": "2026-05-31", "revenue": 15430.44, "units": 1561, "available": True},
    }
    results = iter([unscoped, offline_only])
    monkeypatch.setattr(sellout_service, "get_period_comparison", lambda **kwargs: next(results))

    tool_call = _tool_call_response(
        ("tu_1", "compare_periods", {"current_start": "2026-06-01", "current_end": "2026-06-30", "previous_start": "2026-05-01", "previous_end": "2026-05-31"}),
        ("tu_2", "compare_periods", {"current_start": "2026-06-01", "current_end": "2026-06-30", "previous_start": "2026-05-01", "previous_end": "2026-05-31", "mode": "offline"}),
    )
    _install_fake_client(monkeypatch, [tool_call, _final_answer()])

    result = assistant.ask("compare May to June", history=None, customer="fairprice")

    assert result["chart_series"] == [{"label": "Revenue", "values": [23702.80, 25915.28]}]


def test_compare_periods_single_mode_scoped_call_used_as_is(monkeypatch):
    # A genuinely channel-specific question ("how did offline do vs last
    # week") only ever produces one scoped call -- there's nothing to sum
    # it with, so it's used directly.
    offline_only = {
        "current": {"start": "2026-08-17", "end": "2026-08-23", "revenue": 500.0, "units": 20},
        "previous": {"start": "2026-08-10", "end": "2026-08-16", "revenue": 400.0, "units": 15, "available": True},
    }
    monkeypatch.setattr(sellout_service, "get_period_comparison", lambda **kwargs: offline_only)

    tool_call = _tool_call_response(("tu_1", "compare_periods", {"comparison_type": "wow", "mode": "offline"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer()])

    result = assistant.ask("how did offline do this week", history=None, customer="fairprice")

    assert result["chart_series"] == [{"label": "Revenue", "values": [400.0, 500.0]}]


def test_compare_periods_used_as_short_range_total_shows_single_bar(monkeypatch):
    # Regression (found via live testing): asked for a plain total across a
    # date range, the model correctly used compare_periods as a SQL-summed
    # total calculator (current_start/end only, no comparison_type, no
    # explicit previous_start/end -- exactly the tool description's
    # guidance to avoid hand-adding get_sales_summary rows) -- but the
    # chart still showed an auto-derived "Previous" bar the question never
    # asked about. Without a real comparison request, a short (<=31 day)
    # range shows just the one figure that was actually asked for.
    range_total = {
        "current": {"start": "2026-07-01", "end": "2026-07-19", "revenue": 12000.0, "units": 1200},
        "previous": {"start": "2026-06-01", "end": "2026-06-19", "revenue": 11000.0, "units": 1100, "available": True},
    }
    monkeypatch.setattr(sellout_service, "get_period_comparison", lambda **kwargs: range_total)

    tool_call = _tool_call_response(
        ("tu_1", "compare_periods", {"current_start": "2026-07-01", "current_end": "2026-07-19", "mode": "offline"})
    )
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="Fairprice's offline sales from Jul 1 to Jul 19, 2026 were $12,000.")])

    result = assistant.ask("what was fairprice offline sales from 1 jul 2026 until 19 july 2026", history=None, customer="fairprice")

    assert result["chart_categories"] == ["2026-07-01 to 2026-07-19"]
    assert result["chart_series"] == [{"label": "Revenue", "values": [12000.0]}]


def test_compare_periods_used_as_multi_month_range_total_shows_monthly_breakdown(monkeypatch):
    # Regression (found via live testing): a multi-month range total (e.g.
    # "offline sales from Jan to Jul") showed one flat bar spanning the
    # whole range -- far less useful than a monthly trend. A range over a
    # month long now fetches a fresh monthly breakdown server-side (not via
    # another model tool call, so scope can't drift between two calls).
    range_total = {
        "current": {"start": "2026-01-01", "end": "2026-07-31", "revenue": 123860.85, "units": 12000},
        "previous": {"start": "2025-12-01", "end": "2025-12-31", "revenue": 20000.0, "units": 1800, "available": True},
    }
    monkeypatch.setattr(sellout_service, "get_period_comparison", lambda **kwargs: range_total)

    monthly_summary = {
        "offline": {
            "storeFormats": [],
            "periodTotal": [
                {"period_label": "January 2026", "period_start": "2026-01-01", "revenue": 18496.44, "units": 1711},
                {"period_label": "February 2026", "period_start": "2026-02-05", "revenue": 16041.01, "units": 1466},
            ],
            "periodByFormat": [],
        },
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    captured_kwargs = {}

    def _fake_dashboard_summary(**kwargs):
        captured_kwargs.update(kwargs)
        return monthly_summary

    monkeypatch.setattr(sellout_service, "get_dashboard_summary", _fake_dashboard_summary)

    tool_call = _tool_call_response(
        ("tu_1", "compare_periods", {"current_start": "2026-01-01", "current_end": "2026-07-31", "mode": "offline"})
    )
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="Fairprice's offline sales from Jan 1 to Jul 31, 2026 were $123,860.85.")])

    result = assistant.ask("what was fairprice offline sales from 1 jan 2026 until 31 july 2026", history=None, customer="fairprice")

    assert result["chart_categories"] == ["January 2026", "February 2026"]
    assert result["chart_series"] == [{"label": "Revenue", "values": [18496.44, 16041.01]}]
    # the breakdown fetch reused the exact same scope compare_periods used
    assert captured_kwargs["customer"] == "fairprice"
    assert captured_kwargs["start_date"] == "2026-01-01"
    assert captured_kwargs["end_date"] == "2026-07-31"
    assert captured_kwargs["granularity"] == "month"


def test_multi_month_range_total_falls_back_to_single_bar_if_breakdown_fetch_fails(monkeypatch):
    range_total = {
        "current": {"start": "2026-01-01", "end": "2026-07-31", "revenue": 123860.85, "units": 12000},
        "previous": {"start": "2025-12-01", "end": "2025-12-31", "revenue": 20000.0, "units": 1800, "available": True},
    }
    monkeypatch.setattr(sellout_service, "get_period_comparison", lambda **kwargs: range_total)

    def _boom(**kwargs):
        raise RuntimeError("boom")

    monkeypatch.setattr(sellout_service, "get_dashboard_summary", _boom)

    tool_call = _tool_call_response(
        ("tu_1", "compare_periods", {"current_start": "2026-01-01", "current_end": "2026-07-31", "mode": "offline"})
    )
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="Fairprice's offline sales from Jan 1 to Jul 31, 2026 were $123,860.85.")])

    result = assistant.ask("what was fairprice offline sales from 1 jan 2026 until 31 july 2026", history=None, customer="fairprice")

    assert result["chart_categories"] == ["2026-01-01 to 2026-07-31"]
    assert result["chart_series"] == [{"label": "Revenue", "values": [123860.85]}]


def test_compare_periods_with_no_data_yields_empty_chart(monkeypatch):
    empty = {
        "current": {"start": None, "end": None, "revenue": 0.0, "units": 0.0},
        "previous": {"start": None, "end": None, "revenue": 0.0, "units": 0.0, "available": False},
    }
    monkeypatch.setattr(sellout_service, "get_period_comparison", lambda **kwargs: empty)

    tool_call = _tool_call_response(("tu_1", "compare_periods", {"comparison_type": "wow"}))
    _install_fake_client(
        monkeypatch,
        [
            tool_call,
            _final_answer(grounded=False, data_source="none"),
            _search_answer("Nothing relevant turned up."),
            _final_answer(grounded=False, data_source="none"),
        ],
    )

    result = assistant.ask("how does this week compare", history=None, customer="fairprice")

    assert result["has_chart"] is False


def test_multiple_sales_summary_calls_merge_into_one_comparison_chart(monkeypatch):
    # Regression: answering "compare May vs June" needs a per-format
    # breakdown for each month, which get_sales_summary only gives one
    # month at a time -- so the model calls it twice. The chart must show
    # both months, not just whichever call happened last.
    may = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "May 2026", "period_start": "2026-05-01", "revenue": 23702.80, "units": 2441}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    june = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "June 2026", "period_start": "2026-06-01", "revenue": 25915.28, "units": 2742}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    results = iter([may, june])
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: next(results))

    tool_call = _tool_call_response(
        ("tu_1", "get_sales_summary", {"granularity": "month", "start_date": "2026-05-01", "end_date": "2026-05-31"}),
        ("tu_2", "get_sales_summary", {"granularity": "month", "start_date": "2026-06-01", "end_date": "2026-06-30"}),
    )
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="June beat May.")])

    result = assistant.ask("compare May 2026 to June 2026", history=None, customer="fairprice")

    assert result["has_chart"] is True
    assert result["chart_categories"] == ["May 2026", "June 2026"]
    assert result["chart_series"] == [{"label": "Revenue", "values": [23702.80, 25915.28]}]


def test_sequential_sales_summary_calls_across_turns_also_merge(monkeypatch):
    # Same as above but the two calls happen in separate turns (one tool
    # call per turn) rather than both in one model message.
    may = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "May 2026", "period_start": "2026-05-01", "revenue": 100.0, "units": 1}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    june = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "June 2026", "period_start": "2026-06-01", "revenue": 200.0, "units": 2}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    results = iter([may, june])
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: next(results))

    first_call = _tool_call_response(("tu_1", "get_sales_summary", {"start_date": "2026-05-01", "end_date": "2026-05-31"}))
    second_call = _tool_call_response(("tu_2", "get_sales_summary", {"start_date": "2026-06-01", "end_date": "2026-06-30"}))
    _install_fake_client(monkeypatch, [first_call, second_call, _final_answer()])

    result = assistant.ask("compare May to June", history=None, customer="fairprice")

    assert result["chart_categories"] == ["May 2026", "June 2026"]
    assert result["chart_series"] == [{"label": "Revenue", "values": [100.0, 200.0]}]


def test_compare_periods_wins_over_get_sales_summary_when_both_called(monkeypatch):
    # Regression: the model often calls both tools to answer a comparison
    # question -- compare_periods for the headline total, get_sales_summary
    # for a per-format breakdown in the narrative. The chart must reflect
    # compare_periods' ALL-CHANNEL total (what the answer actually cites),
    # not get_sales_summary's per-channel split, regardless of call order.
    comparison = {
        "current": {"start": "2026-06-01", "end": "2026-06-30", "revenue": 25915.28, "units": 2742},
        "previous": {"start": "2026-05-01", "end": "2026-05-31", "revenue": 23702.80, "units": 2441, "available": True},
    }
    offline_only_summary = {
        "offline": {"storeFormats": [], "periodTotal": [
            {"period_label": "May 2026", "period_start": "2026-05-01", "revenue": 15430.44, "units": 1561},
            {"period_label": "June 2026", "period_start": "2026-06-01", "revenue": 16601.92, "units": 1735},
        ], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_period_comparison", lambda **kwargs: comparison)
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: offline_only_summary)

    # get_sales_summary called SECOND (i.e. it would "win" under a
    # last-call-wins rule) -- compare_periods must still be preferred.
    tool_call = _tool_call_response(
        ("tu_1", "compare_periods", {"current_start": "2026-06-01", "current_end": "2026-06-30", "previous_start": "2026-05-01", "previous_end": "2026-05-31"}),
        ("tu_2", "get_sales_summary", {"start_date": "2026-05-01", "end_date": "2026-06-30"}),
    )
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="June beat May: $25,915.28 vs $23,702.80.")])

    result = assistant.ask("compare May to June", history=None, customer="fairprice")

    assert result["chart_categories"] == ["Previous", "Current"]
    assert result["chart_series"] == [{"label": "Revenue", "values": [23702.80, 25915.28]}]


def test_rank_skus_builds_table_not_chart(monkeypatch):
    df = pd.DataFrame(
        [
            {"sku": "A", "product_name": "Widget", "volume": 10, "value": 100.0, "rank": 1},
            {"sku": "B", "product_name": "Gadget", "volume": 5, "value": 50.0, "rank": 2},
        ]
    )
    monkeypatch.setattr(sellout_service, "get_sku_ranking", lambda **kwargs: df)

    tool_call = _tool_call_response(("tu_1", "rank_skus", {"metric": "value"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer()])

    result = assistant.ask("best sellers?", history=None, customer="fairprice")

    assert result["has_chart"] is False
    assert result["has_table"] is True
    assert result["table_columns"] == ["rank", "product_name", "volume", "value"]
    assert result["table_rows"][0] == ["1", "Widget", "10", "100.0"]


# ---- ask(): inventory tools --------------------------------------------------

def test_inventory_overview_builds_table_not_chart(monkeypatch):
    overview = {
        "customers": [{"customer_id": 1, "customer_name": "fairprice"}],
        "monthly": [{"customer_id": 1, "customer_name": "fairprice", "month": "2026-07-01", "ending_stock": 663}],
        "skus": [
            {"customer_id": 1, "customer_name": "fairprice", "sku": "S1", "product_name": "Adult Pants L",
             "month": "2026-07-01", "ending_stock": 663},
        ],
    }
    monkeypatch.setattr(inventory_service, "get_overview", lambda **kwargs: overview)

    tool_call = _tool_call_response(("tu_1", "get_inventory_overview", {}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="Ending stock is 663 units.")])

    result = assistant.ask("how much stock do we have?", history=None, customer="fairprice")

    assert result["has_chart"] is False
    assert result["has_table"] is True
    assert result["table_columns"] == ["month", "customer", "product", "ending_stock"]
    assert result["table_rows"][0] == ["2026-07-01", "fairprice", "Adult Pants L", "663"]


def test_inventory_overview_omits_customer_filter_when_not_named(monkeypatch):
    # Overview is a cross-customer view by default -- unlike the sales
    # tools, it must NOT silently fall back to the page's current customer.
    captured = {}

    def _fake_overview(**kwargs):
        captured.update(kwargs)
        return {"customers": [], "monthly": [], "skus": []}

    monkeypatch.setattr(inventory_service, "get_overview", _fake_overview)

    tool_call = _tool_call_response(("tu_1", "get_inventory_overview", {}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer(grounded=False, data_source="none"), _search_answer("n/a"), _final_answer(grounded=False, data_source="none")])

    assistant.ask("how much stock do we have across all customers?", history=None, customer="fairprice")

    assert captured["customer_ids"] is None


def test_customer_inventory_resolves_name_and_builds_table(monkeypatch):
    monkeypatch.setattr(inventory_service, "list_customers", lambda: [{"customer_id": 1, "customer_name": "fairprice"}])

    captured = {}

    def _fake_customer_view(**kwargs):
        captured.update(kwargs)
        return {
            "customer": {"customer_id": 1, "customer_name": "fairprice"},
            "threshold": {"target_doh": 30.0, "min_doh": 25.0, "max_doh": 35.0, "is_global_default": True},
            "trend": [],
            "skus": [
                {"sku": "S1", "product_name": "Adult Pants L", "month": "2026-07-01", "ending_stock": 663,
                 "doh": 39.8, "target_doh": 30.0, "doh_status": "above_max"},
            ],
        }

    monkeypatch.setattr(inventory_service, "get_customer_view", _fake_customer_view)

    tool_call = _tool_call_response(("tu_1", "get_customer_inventory", {"customer": "fairprice"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="FairPrice's Adult Pants L is above max at 39.8 days.")])

    result = assistant.ask("what's fairprice's DOH?", history=None, customer="fairprice")

    assert captured["customer_id"] == 1
    assert result["table_columns"] == ["month", "product", "ending_stock", "doh", "target_doh", "status"]
    assert result["table_rows"][0] == ["2026-07-01", "Adult Pants L", "663", "39.8", "30.0", "Above max"]


def test_customer_inventory_unknown_name_is_reported_as_tool_error(monkeypatch):
    monkeypatch.setattr(inventory_service, "list_customers", lambda: [{"customer_id": 1, "customer_name": "fairprice"}])

    tool_call = _tool_call_response(("tu_1", "get_customer_inventory", {"customer": "cold storage"}))
    fake_client = _install_fake_client(
        monkeypatch,
        [tool_call, _final_answer(grounded=False, data_source="none"), _search_answer("n/a"), _final_answer(grounded=False, data_source="none")],
    )

    assistant.ask("what's cold storage's DOH?", history=None, customer="fairprice")

    second_call_contents = fake_client.models.calls[1]["contents"]
    function_response = _dump(second_call_contents[-1])["parts"][0]["function_response"]["response"]
    assert "error" in function_response
    assert "Unknown inventory customer" in function_response["error"]


def test_at_risk_inventory_builds_table_sorted_by_severity(monkeypatch):
    at_risk = {
        "as_of": "2026-07-01",
        "counts": {"below_min": 0, "above_max": 1},
        "items": [
            {"customer_name": "fairprice", "product_name": "Adult Pants L", "doh": 39.8, "target_doh": 30.0,
             "doh_status": "above_max", "days_outside_band": 4.8},
        ],
    }
    monkeypatch.setattr(inventory_service, "get_at_risk", lambda **kwargs: at_risk)

    tool_call = _tool_call_response(("tu_1", "get_at_risk_inventory", {}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="Adult Pants L is overstocked.")])

    result = assistant.ask("what's at risk?", history=None, customer="fairprice")

    assert result["table_columns"] == ["customer", "product", "doh", "target_doh", "status", "days_outside_band"]
    assert result["table_rows"][0] == ["fairprice", "Adult Pants L", "39.8", "30.0", "Above max", "4.8"]


def test_sell_in_plan_builds_table(monkeypatch):
    monkeypatch.setattr(inventory_service, "list_customers", lambda: [{"customer_id": 1, "customer_name": "fairprice"}])
    plan = {
        "customer": {"customer_id": 1, "customer_name": "fairprice"},
        "threshold": {"target_doh": 30.0, "min_doh": 25.0, "max_doh": 35.0},
        "actuals_through": "2026-07-01",
        "actuals_through_by_sku": {},
        "rows": [
            {"month": "2026-08-01", "product_name": "Adult Pants L", "forecast_sell_out": 620.0,
             "recommended_sell_in": 376, "projected_ending_stock": 419.0},
        ],
        "monthly_totals": [],
        "sku_totals": [],
        "skus_without_forecast": [],
    }
    monkeypatch.setattr(inventory_service, "get_sell_in_plan", lambda **kwargs: plan)

    tool_call = _tool_call_response(("tu_1", "get_sell_in_plan", {"customer": "fairprice"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="Recommended sell-in for August is 376 units.")])

    result = assistant.ask("how much should we sell in next month?", history=None, customer="fairprice")

    assert result["table_columns"] == ["month", "product", "forecast_sell_out", "recommended_sell_in", "projected_ending_stock"]
    assert result["table_rows"][0] == ["2026-08-01", "Adult Pants L", "620.0", "376", "419.0"]


def test_list_inventory_skus_lookup(monkeypatch):
    monkeypatch.setattr(inventory_service, "list_skus", lambda: [{"sku": "S1", "product_name": "Adult Pants L"}])

    tool_call = _tool_call_response(("tu_1", "list_inventory_skus", {}))
    fake_client = _install_fake_client(monkeypatch, [tool_call, _final_answer()])

    assistant.ask("what SKUs do we track?", history=None, customer="fairprice")

    second_call_contents = fake_client.models.calls[1]["contents"]
    function_response = _dump(second_call_contents[-1])["parts"][0]["function_response"]["response"]
    assert json.loads(function_response["result"]) == [{"sku": "S1", "product_name": "Adult Pants L"}]


def test_inventory_database_error_is_reported_as_tool_error_not_raised(monkeypatch):
    def _boom(**kwargs):
        raise RuntimeError("connection refused")

    monkeypatch.setattr(inventory_service, "get_at_risk", _boom)

    tool_call = _tool_call_response(("tu_1", "get_at_risk_inventory", {}))
    fake_client = _install_fake_client(
        monkeypatch,
        [tool_call, _final_answer(grounded=False, data_source="none"), _search_answer("n/a"), _final_answer(grounded=False, data_source="none")],
    )

    assistant.ask("what's at risk?", history=None, customer="fairprice")

    second_call_contents = fake_client.models.calls[1]["contents"]
    function_response = _dump(second_call_contents[-1])["parts"][0]["function_response"]["response"]
    assert "Unable to reach the inventory database" in function_response["error"]


def test_a_bigquery_outage_inside_an_inventory_tool_is_also_a_tool_error_not_a_crash(monkeypatch):
    # get_customer_view and get_sell_in_plan both call BigQuery internally (the
    # sell-out/forecast lookups) -- a GoogleAPICallError from there must be
    # caught by the inventory branch's own broad except, not left to propagate
    # past it to the generic sales-tool GoogleAPICallError handler further out,
    # which would give a less specific "Unable to reach BigQuery" answer with
    # no inventory framing.
    def _boom(**kwargs):
        raise ServiceUnavailable("BigQuery is down")

    monkeypatch.setattr(inventory_service, "get_customer_view", _boom)

    tool_call = _tool_call_response(("tu_1", "get_customer_inventory", {"customer": "fairprice"}))
    fake_client = _install_fake_client(
        monkeypatch,
        [tool_call, _final_answer(grounded=False, data_source="none"), _search_answer("n/a"), _final_answer(grounded=False, data_source="none")],
    )

    assistant.ask("what's fairprice's DOH?", history=None, customer="fairprice")

    second_call_contents = fake_client.models.calls[1]["contents"]
    function_response = _dump(second_call_contents[-1])["parts"][0]["function_response"]["response"]
    assert "Unable to reach the inventory database" in function_response["error"]


# ---- ask(): forecast (get_demand_forecast) ------------------------------------

def _forecast_result(**overrides):
    result = {
        "customer": "fairprice",
        "product_name": "Aire Ultra Tape L",
        "channel": "Online and offline combined",
        "unit": "cartons",
        "period": {"start": "2026-09-01", "end": "2026-10-01"},
        "forecast_months_available": {"first": "2026-09-01", "last": "2027-08-01"},
        "forecast_generated_at": "2026-10-04T17:14:40Z",
        "forecast_generated_at_sgt": "5 Oct 2026, 1:14 am",
        "actuals_loaded_at": "2026-10-06T05:05:11Z",
        "actuals_loaded_at_sgt": "6 Oct 2026, 1:05 pm",
        "series": [{
            "customer": "fairprice",
            "product_name": "Aire Ultra Tape L",
            "sku": "13271338",
            "confidence": {"band": "low", "label": "Low", "reason": "limited sales history, check before using", "detail": ""},
            "forecast": [
                {"month": "2026-09-01", "cartons": 52, "low_80": 32, "high_80": 79},
                {"month": "2026-10-01", "cartons": 1052, "low_80": -2, "high_80": 1080},
            ],
            "actuals": [{"month": "2026-08-01", "cartons": 88, "month_to_date": False}],
        }],
        "monthly_totals": [],
        "forecast_total_cartons": 1104,
    }
    return {**result, **overrides}


def _ask_forecast(monkeypatch, answer, result=None):
    seen = {}

    def fake_forecast(**kwargs):
        seen.update(kwargs)
        return result or _forecast_result()

    monkeypatch.setattr(assistant.forecast_summary, "get_demand_forecast", fake_forecast)
    tool_call = _tool_call_response(("tu_1", "get_demand_forecast", {"product": "ultra tape l", "customer": "fairprice"}))
    fake_client = _install_fake_client(monkeypatch, [tool_call, _final_answer(text=answer), _final_answer(text=answer)])
    return assistant.ask("what is the forecast for ultra tape l?", history=None, customer="fairprice"), seen, fake_client


def test_forecast_tool_passes_names_through_and_builds_labelled_table(monkeypatch):
    result, seen, _ = _ask_forecast(monkeypatch, "Forecast for Sep 2026: 52 cartons.")

    assert seen == {"product": "ultra tape l", "customer": "fairprice", "start_month": None, "end_month": None}
    assert result["has_chart"] is False
    assert result["table_columns"] == ["type", "month", "customer", "product", "cartons", "range_80", "confidence"]
    assert result["table_rows"] == [
        ["Actual", "Aug 2026", "Fairprice", "Aire Ultra Tape L (13271338)", "88", "—", "—"],
        ["Forecast", "Sep 2026", "Fairprice", "Aire Ultra Tape L (13271338)", "52", "32 to 79", "Low"],
        ["Forecast", "Oct 2026", "Fairprice", "Aire Ultra Tape L (13271338)", "1,052", "(2) to 1,080", "Low"],
        ["Forecast total", "Sep 2026 – Oct 2026", "Fairprice", "Aire Ultra Tape L", "1,104", "—", "—"],
    ]
    assert result["table_note"] == (
        "Forecast run: 5 Oct 2026, 1:14 am SGT · Actuals loaded: 6 Oct 2026, 1:05 pm SGT"
        " · Cartons; online and offline combined, as on the Forecast page"
    )


def test_forecast_month_to_date_actual_is_labelled_as_such(monkeypatch):
    forecast = _forecast_result()
    forecast["series"][0]["actuals"] = [{"month": "2026-10-01", "cartons": 20, "month_to_date": True}]

    result, _, _ = _ask_forecast(monkeypatch, "Forecast for Sep 2026: 52 cartons.", result=forecast)

    assert result["table_rows"][0][0] == "Actual (month to date)"


def test_forecast_with_many_series_shows_only_month_totals(monkeypatch):
    series = [
        {**_forecast_result()["series"][0], "product_name": f"SKU {i}", "sku": str(i)}
        for i in range(assistant._MAX_FORECAST_SERIES_IN_TABLE + 1)
    ]
    totals = [{"month": "2026-09-01", "forecast_cartons": 520, "actual_cartons": None, "actual_month_to_date": False}]
    forecast = _forecast_result(series=series, monthly_totals=totals, product_name=None, forecast_total_cartons=None)

    result, _, _ = _ask_forecast(monkeypatch, "Forecast for Sep 2026: 520 cartons.", result=forecast)

    assert result["table_rows"] == [["Forecast total", "Sep 2026", "Fairprice", "All SKUs", "520", "—", "—"]]


def test_forecast_tool_bad_arguments_are_a_tool_error_not_a_crash(monkeypatch):
    def unknown_customer(**kwargs):
        raise ValueError("No forecast for customer 'nhg'.")

    monkeypatch.setattr(assistant.forecast_summary, "get_demand_forecast", unknown_customer)
    tool_call = _tool_call_response(("tu_1", "get_demand_forecast", {"customer": "nhg"}))
    fake_client = _install_fake_client(monkeypatch, [tool_call, _final_answer(text="There is no forecast for NHG yet.")])

    result = assistant.ask("forecast for nhg?", history=None, customer="fairprice")

    tool_turn = _dump(fake_client.models.calls[1]["contents"][-1])
    assert "Invalid arguments: No forecast for customer 'nhg'." in json.dumps(tool_turn)
    assert result["has_table"] is False


def test_forecast_table_wins_over_a_sales_lookup_made_alongside_it():
    chart = assistant._build_chart(
        [
            ("get_sales_summary", {}, {"offline": {"periodTotal": [{"period_label": "Aug 2026", "revenue": 712.3}]}}),
            ("get_demand_forecast", {}, _forecast_result()),
        ],
        "fairprice",
    )

    assert chart["has_table"] is True
    assert chart["table_rows"][1][0] == "Forecast"


def test_forecast_headline_cartons_matching_the_table_is_kept(monkeypatch):
    result, _, fake_client = _ask_forecast(monkeypatch, "Forecast for Oct 2026: 1,052 cartons (Confidence: Low).")

    assert len(fake_client.models.calls) == 2
    assert result["has_table"] is True


def test_forecast_headline_cartons_not_in_the_table_hides_the_table(monkeypatch):
    result, _, fake_client = _ask_forecast(monkeypatch, "Forecast for Sep 2026: 60 cartons.")

    # One correction retry, still wrong: the table that disagrees is left out.
    assert len(fake_client.models.calls) == 3
    assert result["has_table"] is False


def _pptx_text(pptx_bytes):
    from pptx import Presentation

    texts = []
    for slide in Presentation(io.BytesIO(pptx_bytes)).slides:
        for shape in slide.shapes:
            if shape.has_table:
                texts.extend(" ".join(cell.text for cell in row.cells) for row in shape.table.rows)
            elif shape.has_text_frame:
                texts.append(shape.text_frame.text)
    return texts


def test_forecast_export_pdf_and_pptx_include_the_table_and_note():
    chart = assistant._forecast_chart([("get_demand_forecast", {}, _forecast_result())])

    pdf_bytes, _, _ = assistant.generate_export_file("pdf", "Forecast for Sep 2026: 52 cartons.", chart)
    pptx_bytes, _, _ = assistant.generate_export_file("pptx", "Forecast for Sep 2026: 52 cartons.", chart)

    assert pdf_bytes[:4] == b"%PDF"
    texts = _pptx_text(pptx_bytes)
    assert any(text.startswith("Forecast Sep 2026") and "52" in text for text in texts)
    assert any(text.startswith("Actual Aug 2026") and "88" in text for text in texts)
    assert any(text.startswith("Forecast run: 5 Oct 2026, 1:14 am SGT") for text in texts)


def test_long_table_continues_on_further_slides_with_its_header():
    from pptx import Presentation

    rows = [[str(i), "x"] for i in range(assistant._MAX_TABLE_ROWS_PER_SLIDE + 3)]
    chart = assistant._table_chart(["n", "v"], rows)

    pptx_bytes, _, _ = assistant.generate_export_file("pptx", "Some answer.", chart)

    tables = [shape.table for slide in Presentation(io.BytesIO(pptx_bytes)).slides for shape in slide.shapes if shape.has_table]
    assert [len(table.rows) for table in tables] == [assistant._MAX_TABLE_ROWS_PER_SLIDE + 1, 4]
    assert all(table.cell(0, 0).text == "n" for table in tables)


def test_export_endpoint_passes_the_table_note_to_the_file(monkeypatch):
    from fastapi.testclient import TestClient
    from app.main import app

    seen = {}

    def fake_export(export_format, answer, chart, chart_style, trend_values):
        seen["chart"] = chart
        return b"%PDF", "application/pdf", "aireos-report.pdf"

    monkeypatch.setattr(assistant, "generate_export_file", fake_export)

    response = TestClient(app).post("/api/assistant/export", json={
        "format": "pdf", "answer": "a", "has_table": True,
        "table_columns": ["type"], "table_rows": [["Forecast"]], "table_note": "Forecast run: 5 Oct 2026",
    })

    assert response.status_code == 200
    assert seen["chart"]["table_note"] == "Forecast run: 5 Oct 2026"


def _priced_forecast_result():
    forecast = _forecast_result()
    forecast["series"][0]["forecast"][0]["revenue"] = 441.79
    forecast["series"][0]["forecast"][1]["revenue"] = 8937.81
    forecast["series"][0]["actuals"][0]["revenue"] = 712.3
    forecast["forecast_total_revenue"] = 9379.6
    return forecast


def _ask_forecast_revenue(monkeypatch, answer):
    monkeypatch.setattr(assistant.forecast_summary, "get_demand_forecast", lambda **kwargs: _priced_forecast_result())
    tool_call = _tool_call_response(
        ("tu_1", "get_demand_forecast", {"product": "ultra tape l", "customer": "fairprice", "include_revenue": True})
    )
    fake_client = _install_fake_client(monkeypatch, [tool_call, _final_answer(text=answer), _final_answer(text=answer)])
    return assistant.ask("forecast revenue for ultra tape l?", history=None, customer="fairprice"), fake_client


def test_forecast_revenue_question_shows_revenue_column_instead_of_range(monkeypatch):
    result, _ = _ask_forecast_revenue(monkeypatch, "Forecast revenue for Sep 2026 is $441.79.")

    assert result["table_columns"] == ["type", "month", "customer", "product", "cartons", "revenue", "confidence"]
    assert result["table_rows"] == [
        ["Actual", "Aug 2026", "Fairprice", "Aire Ultra Tape L (13271338)", "88", "$712.30", "—"],
        ["Forecast", "Sep 2026", "Fairprice", "Aire Ultra Tape L (13271338)", "52", "$441.79", "Low"],
        ["Forecast", "Oct 2026", "Fairprice", "Aire Ultra Tape L (13271338)", "1,052", "$8,937.81", "Low"],
        ["Forecast total", "Sep 2026 – Oct 2026", "Fairprice", "Aire Ultra Tape L", "1,104", "$9,379.60", "—"],
    ]
    assert "realised price per carton" in result["table_note"]


def test_forecast_revenue_headline_matching_the_tool_is_kept(monkeypatch):
    result, fake_client = _ask_forecast_revenue(monkeypatch, "Forecast revenue for Sep–Oct 2026 is $9,379.60.")

    assert len(fake_client.models.calls) == 2
    assert result["has_table"] is True


def test_forecast_revenue_headline_worked_out_by_the_model_hides_the_table(monkeypatch):
    result, fake_client = _ask_forecast_revenue(monkeypatch, "Forecast revenue for Sep 2026 is $450.00.")

    assert len(fake_client.models.calls) == 3
    assert result["has_table"] is False


def test_cartons_question_keeps_the_range_column(monkeypatch):
    result, _, _ = _ask_forecast(monkeypatch, "Forecast for Sep 2026: 52 cartons.")

    assert result["table_columns"][5] == "range_80"
    assert "realised price" not in result["table_note"]


# ---- Customer nicknames -------------------------------------------------------

@pytest.mark.parametrize("typed", ["NTUC", "ntuc fairprice", "FP", "Fair Price", "FairPrice"])
def test_customer_nicknames_resolve_to_the_stored_customer(typed):
    assert assistant._customer_slug(typed) == "fairprice"


def test_unknown_names_pass_through_unchanged():
    assert assistant._customer_slug("Sheng Shiong") == "sheng_shiong"
    assert assistant._customer_slug("Giant") == "giant"


def test_forecast_tool_gets_the_resolved_customer_not_the_nickname(monkeypatch):
    seen = {}

    def fake_forecast(**kwargs):
        seen.update(kwargs)
        return _forecast_result()

    monkeypatch.setattr(assistant.forecast_summary, "get_demand_forecast", fake_forecast)
    tool_call = _tool_call_response(("tu_1", "get_demand_forecast", {"customer": "NTUC"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="Forecast for Sep 2026: 52 cartons.")])

    assistant.ask("forecast for ntuc?", history=None, customer="fairprice")

    assert seen["customer"] == "fairprice"


def test_sales_tools_also_get_the_resolved_customer(monkeypatch):
    seen = {}

    def fake_summary(**kwargs):
        seen.update(kwargs)
        return {"offline": {"periodTotal": []}, "online": {"periodTotal": []}}

    monkeypatch.setattr(sellout_service, "get_dashboard_summary", fake_summary)
    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"customer": "NTUC", "granularity": "month"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="No sales yet.")])

    assistant.ask("ntuc sales?", history=None, customer="fairprice")

    assert seen["customer"] == "fairprice"


# ---- ask(): search fallback --------------------------------------------------
#
# Confirmed live against the real API, two separate constraints: (1) Gemini
# rejects mixing function-declaration tools with search tools in one
# request, and (2) structured output ("controlled generation") can't be
# combined with the search tool either. So a search fallback is genuinely
# three calls in the worst case: business tools (schema, no search) ->
# search (plain text, no schema) -> reshape (schema, no tools) -- the
# fixtures below model that with a plain-text fake for the search step.

def _search_answer(text):
    """The search-tool call returns plain text, not the JSON schema --
    unlike _final_answer(), which represents a schema-constrained call."""
    return FakeResponse(text=text, content={"role": "model", "parts": [{"text": text}]})


def test_search_fallback_used_when_business_tools_cannot_answer(monkeypatch):
    business_attempt = _final_answer(grounded=False, data_source="none", text="I don't have that in our sales data.")
    search_attempt = _search_answer("Industry-wide, apparel sales grew 4% this quarter.")
    reshaped = _final_answer(grounded=True, data_source="market_data", text="Industry-wide, apparel sales grew 4% this quarter.")
    fake_client = _install_fake_client(monkeypatch, [business_attempt, search_attempt, reshaped])

    result = assistant.ask("how is the broader apparel market doing?", history=None, customer="fairprice")

    assert result["grounded"] is True
    assert result["data_source"] == "market_data"
    assert "apparel sales grew" in result["answer"]
    # 1st call: business tools + schema. 2nd: search tool, no schema.
    # 3rd: no tools, schema only (reshaping the search text).
    assert fake_client.models.calls[0]["config"].tools == assistant._BUSINESS_TOOLS
    assert fake_client.models.calls[0]["config"].response_json_schema is not None
    assert fake_client.models.calls[1]["config"].tools == assistant._SEARCH_TOOLS
    assert fake_client.models.calls[1]["config"].response_json_schema is None
    assert fake_client.models.calls[2]["config"].tools is None
    assert fake_client.models.calls[2]["config"].response_json_schema is not None


def test_search_fallback_not_attempted_when_business_tools_succeed(monkeypatch):
    fake_client = _install_fake_client(monkeypatch, [_final_answer(grounded=True)])

    assistant.ask("how did we do last week?", history=None, customer="fairprice")

    assert len(fake_client.models.calls) == 1


def test_search_fallback_also_ungrounded_keeps_original_answer(monkeypatch):
    business_attempt = _final_answer(grounded=False, data_source="none", text="No internal data for that.")
    search_attempt = _search_answer("Nothing relevant turned up.")
    reshaped = _final_answer(grounded=False, data_source="none", text="Search didn't find anything either.")
    _install_fake_client(monkeypatch, [business_attempt, search_attempt, reshaped])

    result = assistant.ask("something genuinely unanswerable", history=None, customer="fairprice")

    assert result["grounded"] is False
    assert result["answer"] == "No internal data for that."


def test_search_fallback_refusal_keeps_original_ungrounded_answer(monkeypatch):
    business_attempt = _final_answer(grounded=False, data_source="none", text="No internal data for that.")
    search_refusal = FakeResponse(finish_reason="SAFETY", content={"role": "model", "parts": []})
    _install_fake_client(monkeypatch, [business_attempt, search_refusal])

    result = assistant.ask("something disallowed to search for", history=None, customer="fairprice")

    assert result["grounded"] is False
    assert result["answer"] == "No internal data for that."


def test_search_fallback_reshape_refusal_keeps_original_ungrounded_answer(monkeypatch):
    # The search step itself succeeds, but the follow-up reshape-into-JSON
    # call is refused -- still falls back to the original answer rather
    # than crashing on an unparseable/absent .text.
    business_attempt = _final_answer(grounded=False, data_source="none", text="No internal data for that.")
    search_attempt = _search_answer("Some search result text.")
    reshape_refusal = FakeResponse(finish_reason="SAFETY", content={"role": "model", "parts": []})
    _install_fake_client(monkeypatch, [business_attempt, search_attempt, reshape_refusal])

    result = assistant.ask("something searchable but unshapeable", history=None, customer="fairprice")

    assert result["grounded"] is False
    assert result["answer"] == "No internal data for that."


# ---- ask(): numeric grounding safeguard --------------------------------------
#
# Belt-and-suspenders on top of chart-building: the chart can never be
# wrong (built purely from tool data), but the model can occasionally slip
# on the headline dollar figure in the *text* even with every input number
# correct (observed live, unreproducible). These check that a mismatched
# headline triggers exactly one correction retry, and that retry is only
# trusted if it verifies too.

def test_mismatched_headline_figure_triggers_correction_retry(monkeypatch):
    summary = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-01", "revenue": 5473.15, "units": 388}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month", "mode": "offline"}))
    wrong_answer = _final_answer(text="Fairprice's offline sales in August 2026 were $5,123.45.")
    corrected_answer = _final_answer(text="Fairprice's offline sales in August 2026 were $5,473.15.")
    fake_client = _install_fake_client(monkeypatch, [tool_call, wrong_answer, corrected_answer])

    result = assistant.ask("what were offline sales in august", history=None, customer="fairprice")

    assert result["answer"] == "Fairprice's offline sales in August 2026 were $5,473.15."
    assert len(fake_client.models.calls) == 3  # tool call, wrong answer, correction retry
    correction_message = fake_client.models.calls[2]["contents"][-1]
    assert "5,123.45" in correction_message["parts"][0]["text"]
    assert "5,473.15" in correction_message["parts"][0]["text"]


def test_matching_headline_figure_does_not_trigger_retry(monkeypatch):
    summary = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-01", "revenue": 5473.15, "units": 388}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month", "mode": "offline"}))
    fake_client = _install_fake_client(
        monkeypatch, [tool_call, _final_answer(text="Fairprice's offline sales in August 2026 were $5,473.15.")]
    )

    assistant.ask("what were offline sales in august", history=None, customer="fairprice")

    assert len(fake_client.models.calls) == 2  # no correction retry


def test_correction_retry_still_wrong_keeps_original_answer(monkeypatch):
    summary = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-01", "revenue": 5473.15, "units": 388}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month", "mode": "offline"}))
    wrong_answer = _final_answer(text="Offline sales in August 2026 were $5,123.45.")
    still_wrong_answer = _final_answer(text="Offline sales in August 2026 were $5,200.00.")
    _install_fake_client(monkeypatch, [tool_call, wrong_answer, still_wrong_answer])

    result = assistant.ask("what were offline sales in august", history=None, customer="fairprice")

    # the retry didn't verify either -- keep the original rather than trust
    # a second unverified guess just because it's the most recent one, and
    # don't show a chart that disagrees with the figure in the text
    assert result["answer"].startswith("Offline sales in August 2026 were $5,123.45.")
    assert "verify" not in result["answer"]
    assert result["has_chart"] is False


def test_market_data_answers_are_not_checked_against_tool_figures(monkeypatch):
    business_attempt = _final_answer(grounded=False, data_source="none", text="No internal data for that.")
    search_attempt = _search_answer("Competitor launched a $999.00 product.")
    reshaped = _final_answer(grounded=True, data_source="market_data", text="A competitor launched a $999.00 product.")
    fake_client = _install_fake_client(monkeypatch, [business_attempt, search_attempt, reshaped])

    result = assistant.ask("any competitor news?", history=None, customer="fairprice")

    assert result["answer"] == "A competitor launched a $999.00 product."
    assert len(fake_client.models.calls) == 3  # no extra correction-retry call


def test_analysis_answers_are_not_checked_against_tool_figures(monkeypatch):
    # data_source="analysis" (reasons/suggestions, not a data lookup) is
    # intentionally ungrounded -- any figure it mentions in its own
    # reasoning shouldn't trigger the numeric-verification safeguard, which
    # only makes sense for aire_data answers claiming to be tool-verified.
    fake_client = _install_fake_client(
        monkeypatch,
        [_final_answer(grounded=False, data_source="analysis", text="Revenue likely dropped due to fewer promotions -- consider running one in $500 increments.")],
    )

    result = assistant.ask("why did revenue drop and what should we do?", history=None, customer="fairprice")

    assert result["data_source"] == "analysis"
    assert len(fake_client.models.calls) == 1  # no correction retry attempted


def test_answer_with_no_dollar_figure_is_not_checked(monkeypatch):
    tool_call = _tool_call_response(("tu_1", "list_customers", {}))
    monkeypatch.setattr(sellout_service, "get_customer_options", lambda: [{"value": "fairprice", "label": "Fairprice"}])
    fake_client = _install_fake_client(
        monkeypatch, [tool_call, _final_answer(text="We currently have data for Fairprice.")]
    )

    result = assistant.ask("which customers do we have?", history=None, customer="fairprice")

    assert result["answer"] == "We currently have data for Fairprice."
    assert len(fake_client.models.calls) == 2  # no correction retry attempted


def test_correction_retry_refusal_keeps_original_answer(monkeypatch):
    summary = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-01", "revenue": 5473.15, "units": 388}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month", "mode": "offline"}))
    wrong_answer = _final_answer(text="Offline sales in August 2026 were $5,123.45.")
    retry_refusal = FakeResponse(finish_reason="SAFETY", content={"role": "model", "parts": []})
    _install_fake_client(monkeypatch, [tool_call, wrong_answer, retry_refusal])

    result = assistant.ask("what were offline sales in august", history=None, customer="fairprice")

    assert result["answer"].startswith("Offline sales in August 2026 were $5,123.45.")
    assert result["has_chart"] is False


# ---- ask(): PDF/PPTX export --------------------------------------------------
#
# File bytes are built entirely from the already-correct `chart` dict, never
# round-tripped through the model -- same reasoning as chart/table numbers
# never being. These check the dispatch/degradation logic; byte-level
# PDF/PPTX correctness was spot-checked manually (SimpleDocTemplate/
# Presentation round-trips verified directly against the installed
# reportlab/python-pptx before writing the production code).

def test_export_format_none_produces_no_download(monkeypatch):
    _install_fake_client(monkeypatch, [_final_answer(export_format="none")])

    result = assistant.ask("how did we do last week?", history=None, customer="fairprice")

    assert result["has_download"] is False
    assert result["download_filename"] is None
    assert result["download_base64"] is None


def test_export_pdf_with_chart_produces_real_pdf_bytes(monkeypatch):
    summary = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-01", "revenue": 5473.15, "units": 388}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month", "mode": "offline"}))
    _install_fake_client(
        monkeypatch,
        [tool_call, _final_answer(text="Offline sales in August 2026 were $5,473.15.", export_format="pdf")],
    )

    result = assistant.ask("export this as pdf", history=None, customer="fairprice")

    assert result["has_download"] is True
    assert result["download_filename"] == "aireos-report.pdf"
    assert result["download_mime_type"] == "application/pdf"
    decoded = base64.b64decode(result["download_base64"])
    assert decoded[:4] == b"%PDF"


def test_export_pptx_with_chart_produces_real_pptx_bytes(monkeypatch):
    summary = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-01", "revenue": 5473.15, "units": 388}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month", "mode": "offline"}))
    _install_fake_client(
        monkeypatch,
        [tool_call, _final_answer(text="Offline sales in August 2026 were $5,473.15.", export_format="pptx")],
    )

    result = assistant.ask("make me a slide for this", history=None, customer="fairprice")

    assert result["has_download"] is True
    assert result["download_filename"] == "aireos-report.pptx"
    assert result["download_mime_type"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    decoded = base64.b64decode(result["download_base64"])
    assert decoded[:2] == b"PK"  # pptx is a zip archive


def test_export_excel_with_chart_produces_real_xlsx_bytes(monkeypatch):
    summary = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-01", "revenue": 5473.15, "units": 388}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month", "mode": "offline"}))
    _install_fake_client(
        monkeypatch,
        [tool_call, _final_answer(text="Offline sales in August 2026 were $5,473.15.", export_format="excel")],
    )

    result = assistant.ask("export this as excel", history=None, customer="fairprice")

    assert result["has_download"] is True
    assert result["download_filename"] == "aireos-report.xlsx"
    assert result["download_mime_type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    decoded = base64.b64decode(result["download_base64"])
    assert decoded[:2] == b"PK"  # xlsx is a zip archive


def test_export_with_table_instead_of_chart_still_produces_a_file(monkeypatch):
    df = pd.DataFrame([{"sku": "A", "product_name": "Widget", "volume": 10, "value": 100.0, "rank": 1}])
    monkeypatch.setattr(sellout_service, "get_sku_ranking", lambda **kwargs: df)

    tool_call = _tool_call_response(("tu_1", "rank_skus", {"metric": "value"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="Widget is the top seller.", export_format="pdf")])

    result = assistant.ask("export our best sellers as pdf", history=None, customer="fairprice")

    assert result["has_download"] is True
    decoded = base64.b64decode(result["download_base64"])
    assert decoded[:4] == b"%PDF"


def test_export_with_no_chart_or_table_still_produces_text_only_file(monkeypatch):
    # A bare "export this as PDF" with nothing fetched this turn -- degrades
    # to a text-only file rather than failing, per the plan's resolution for
    # when the model doesn't re-fetch data on an export follow-up.
    _install_fake_client(monkeypatch, [_final_answer(text="We currently have data for Fairprice.", export_format="pdf")])

    result = assistant.ask("export that as pdf", history=None, customer="fairprice")

    assert result["has_download"] is True
    decoded = base64.b64decode(result["download_base64"])
    assert decoded[:4] == b"%PDF"


def test_split_into_points_requires_at_least_two_bullets():
    assert assistant._split_into_points("Just one plain sentence.") is None
    assert assistant._split_into_points("- Only one bullet") is None
    assert assistant._split_into_points("- First point\n- Second point\n- Third point") == [
        "First point", "Second point", "Third point",
    ]
    assert assistant._split_into_points("Some intro text\n- Point A\n- Point B") == ["Point A", "Point B"]


def test_export_pptx_renders_multiple_insights_as_bulleted_paragraphs(monkeypatch):
    from pptx import Presentation
    import io as _io

    analysis_text = (
        "- Revenue dipped in Week 21, coinciding with Widget A's volume dropping 40%.\n"
        "- The recovery in Week 22 lines up with restocking of Widget A.\n"
        "- Offline sales consistently outpaced online across all four weeks."
    )
    _install_fake_client(
        monkeypatch,
        [_final_answer(text=analysis_text, grounded=False, data_source="analysis", export_format="pptx")],
    )

    result = assistant.ask("give me insights on possible reasons for the recent sales trend", history=None, customer="fairprice")

    decoded = base64.b64decode(result["download_base64"])
    prs = Presentation(_io.BytesIO(decoded))
    title_slide = prs.slides[0]
    body_paragraphs = [p.text for p in title_slide.placeholders[1].text_frame.paragraphs]
    assert body_paragraphs == [
        "Revenue dipped in Week 21, coinciding with Widget A's volume dropping 40%.",
        "The recovery in Week 22 lines up with restocking of Widget A.",
        "Offline sales consistently outpaced online across all four weeks.",
    ]


def test_export_pptx_shrinks_font_and_enables_autofit_for_long_bullets(monkeypatch):
    # Regression for the overflow bug: PowerPoint's default 28pt placeholder
    # font overflowed the slide with just 3-4 full-sentence bullets (seen
    # live, screenshot showed text cut off past the slide edge).
    from pptx import Presentation
    from pptx.enum.text import MSO_AUTO_SIZE
    import io as _io

    analysis_text = (
        "- Sales peaked in April, potentially due to successful promotional campaigns or strong seasonal demand.\n"
        "- A notable dip occurred in May, which could be a natural follow-up after an active promotional period.\n"
        "- The significant decline in August is likely influenced by the Singapore National Day public holiday."
    )
    _install_fake_client(
        monkeypatch,
        [_final_answer(text=analysis_text, grounded=False, data_source="analysis", export_format="pptx")],
    )

    result = assistant.ask("give me insights on the sales trend", history=None, customer="fairprice")

    decoded = base64.b64decode(result["download_base64"])
    prs = Presentation(_io.BytesIO(decoded))
    tf = prs.slides[0].placeholders[1].text_frame
    assert tf.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for paragraph in tf.paragraphs:
        assert paragraph.font.size.pt <= 18  # well under the 28pt default that overflowed


def test_export_pptx_splits_many_bullets_across_multiple_slides(monkeypatch):
    from pptx import Presentation
    import io as _io

    analysis_text = "\n".join(f"- Insight number {i}, a genuinely distinct data-backed point." for i in range(1, 7))
    _install_fake_client(
        monkeypatch,
        [_final_answer(text=analysis_text, grounded=False, data_source="analysis", export_format="pptx")],
    )

    result = assistant.ask("give me insights on the sales trend", history=None, customer="fairprice")

    decoded = base64.b64decode(result["download_base64"])
    prs = Presentation(_io.BytesIO(decoded))
    # 6 bullets, capped at 4/slide -> 2 insight slides, then the chart slide
    assert prs.slides[0].shapes.title.text == "AireOS Report (1/2)"
    assert prs.slides[1].shapes.title.text == "AireOS Report (2/2)"
    slide0_bullets = [p.text for p in prs.slides[0].placeholders[1].text_frame.paragraphs]
    slide1_bullets = [p.text for p in prs.slides[1].placeholders[1].text_frame.paragraphs]
    assert len(slide0_bullets) == 4
    assert len(slide1_bullets) == 2
    assert slide0_bullets + slide1_bullets == [f"Insight number {i}, a genuinely distinct data-backed point." for i in range(1, 7)]


def test_export_pptx_single_sentence_answer_still_uses_plain_title_slide(monkeypatch):
    from pptx import Presentation
    import io as _io

    _install_fake_client(
        monkeypatch,
        [_final_answer(text="Revenue was $1,000 last week.", export_format="pptx")],
    )

    result = assistant.ask("how did we do last week?", history=None, customer="fairprice")

    decoded = base64.b64decode(result["download_base64"])
    prs = Presentation(_io.BytesIO(decoded))
    assert prs.slides[0].placeholders[1].text == "Revenue was $1,000 last week."


def test_export_pdf_renders_multiple_insights_as_bullets(monkeypatch):
    analysis_text = "- First data-backed insight.\n- Second data-backed insight."
    _install_fake_client(
        monkeypatch,
        [_final_answer(text=analysis_text, grounded=False, data_source="analysis", export_format="pdf")],
    )

    result = assistant.ask("give me insights on the sales trend", history=None, customer="fairprice")

    decoded = base64.b64decode(result["download_base64"])
    assert decoded[:4] == b"%PDF"


def test_export_generation_failure_degrades_to_no_download(monkeypatch):
    monkeypatch.setattr(assistant, "_build_pdf", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom")))
    _install_fake_client(monkeypatch, [_final_answer(export_format="pdf")])

    result = assistant.ask("export this as pdf", history=None, customer="fairprice")

    assert result["has_download"] is False
    assert result["answer"] == "Revenue was $1,000 last week."  # the core answer still comes through


# ---- ask(): refusal / loop-limit --------------------------------------------

def test_refusal_returns_safe_fallback_instead_of_parsing_content(monkeypatch):
    blocked = FakeResponse(finish_reason="SAFETY", content={"role": "model", "parts": []})
    _install_fake_client(monkeypatch, [blocked])

    result = assistant.ask("something disallowed", history=None, customer="fairprice")

    assert result["grounded"] is False
    assert result["data_source"] == "none"
    assert result["has_chart"] is False


def test_a_loop_that_never_settles_falls_back_to_the_web_search(monkeypatch):
    tool_call = _tool_call_response(("tu_1", "list_customers", {}))
    monkeypatch.setattr(sellout_service, "get_customer_options", lambda: [])
    responses = [tool_call] * (assistant.MAX_TOOL_ITERATIONS + 1) + [
        _search_answer("Competitor news."),
        _final_answer(grounded=True, data_source="market_data", text="A competitor launched a product."),
    ]
    fake_client = _install_fake_client(monkeypatch, responses)

    result = assistant.ask("competitor news?", history=None, customer="fairprice")

    assert result["answer"] == "A competitor launched a product."
    assert result["data_source"] == "market_data"
    # the dangling tool call from the failed last turn is not replayed to the search request
    search_contents = fake_client.models.calls[-2]["contents"]
    assert _dump(search_contents[-2])["parts"][0].get("function_response") is not None


def test_a_loop_that_never_settles_and_a_search_that_finds_nothing_gives_a_plain_answer(monkeypatch):
    tool_call = _tool_call_response(("tu_1", "list_customers", {}))
    monkeypatch.setattr(sellout_service, "get_customer_options", lambda: [])
    empty_search = FakeResponse(text=None, content={"role": "model", "parts": []})
    responses = [tool_call] * (assistant.MAX_TOOL_ITERATIONS + 1) + [empty_search]
    _install_fake_client(monkeypatch, responses)

    result = assistant.ask("something unanswerable", history=None, customer="fairprice")

    assert result["answer"] == "I couldn't find an answer to that."
    assert result["has_chart"] is False and result["has_table"] is False


def test_after_max_iterations_the_model_gets_one_last_turn_without_tools(monkeypatch):
    tool_call = _tool_call_response(("tu_1", "list_customers", {}))
    monkeypatch.setattr(sellout_service, "get_customer_options", lambda: [])
    responses = [tool_call] * assistant.MAX_TOOL_ITERATIONS + [_final_answer(text="Here is what I found.", data_source="none")]
    fake_client = _install_fake_client(monkeypatch, responses)

    result = assistant.ask("a question the model keeps researching", history=None, customer="fairprice")

    assert result["answer"] == "Here is what I found."
    last_config = fake_client.models.calls[-1]["config"]
    assert last_config.tool_config.function_calling_config.mode == "NONE"
    assert len(fake_client.models.calls) == assistant.MAX_TOOL_ITERATIONS + 1


# ---- get_client() config error ----------------------------------------------

def test_get_client_raises_when_gcp_project_missing(monkeypatch):
    monkeypatch.setattr(assistant, "_client", None)
    monkeypatch.setattr(assistant, "GCP_PROJECT_ID", None)

    with pytest.raises(assistant.AssistantConfigError):
        assistant.get_client()


# ---- chart_style: pure helpers ----------------------------------------------

def test_sanitize_color_accepts_default_and_valid_hex():
    assert assistant._sanitize_color("default") == "default"
    assert assistant._sanitize_color("#DC2626") == "#DC2626"


def test_sanitize_color_rejects_invalid_values():
    assert assistant._sanitize_color("red") == "default"
    assert assistant._sanitize_color("#GGGGGG") == "default"
    assert assistant._sanitize_color(None) == "default"
    assert assistant._sanitize_color("") == "default"


def test_linear_trend_on_perfectly_linear_series():
    assert assistant._linear_trend([10.0, 20.0, 30.0]) == [10.0, 20.0, 30.0]


def test_linear_trend_fits_two_points_exactly():
    assert assistant._linear_trend([1000.0, 1500.0]) == [1000.0, 1500.0]


def test_linear_trend_returns_empty_for_fewer_than_two_points():
    assert assistant._linear_trend([]) == []
    assert assistant._linear_trend([100.0]) == []


def test_resolve_last_chart_falls_back_when_no_new_chart_data():
    last_chart = {
        "has_chart": True, "chart_type": "bar", "chart_categories": ["May"],
        "chart_series": [{"label": "Revenue", "values": [500.0]}],
        "has_table": False, "table_columns": [], "table_rows": [],
    }
    assert assistant._resolve_last_chart(assistant._empty_chart(), last_chart) == last_chart


def test_resolve_last_chart_prefers_fresh_data_over_stale():
    fresh_chart = {
        "has_chart": True, "chart_type": "bar", "chart_categories": ["June"],
        "chart_series": [{"label": "Revenue", "values": [700.0]}],
        "has_table": False, "table_columns": [], "table_rows": [],
    }
    stale_chart = {**fresh_chart, "chart_categories": ["May"]}
    assert assistant._resolve_last_chart(fresh_chart, stale_chart) == fresh_chart


def test_resolve_last_chart_stays_empty_when_neither_has_data():
    assert assistant._resolve_last_chart(assistant._empty_chart(), None) == assistant._empty_chart()
    assert assistant._resolve_last_chart(assistant._empty_chart(), {}) == assistant._empty_chart()


def test_system_prompt_states_the_current_chart_style():
    style = {**assistant._DEFAULT_CHART_STYLE, "color": "#DC2626", "show_value_labels": True}
    prompt = assistant._system_prompt(style)
    assert "color=#DC2626" in prompt
    assert "show_value_labels=True" in prompt
    assert "show_trend_line=False" in prompt
    assert "value_label_format=auto" in prompt
    assert "show_gridlines=True" in prompt
    assert "chart_type=auto" in prompt


def test_sanitize_chart_style_falls_back_to_defaults_for_bad_enums():
    result = assistant._sanitize_chart_style({
        "color": "not-a-color",
        "show_value_labels": "yes",
        "value_label_format": "scientific-notation",
        "show_trend_line": 1,
        "show_gridlines": "nope",
        "chart_type": "pie",
    })
    assert result["color"] == "default"
    assert result["show_value_labels"] is True
    assert result["value_label_format"] == "auto"
    assert result["show_trend_line"] is True
    assert result["show_gridlines"] is True
    assert result["chart_type"] == "auto"


def test_sanitize_chart_style_handles_missing_input():
    assert assistant._sanitize_chart_style(None) == assistant._DEFAULT_CHART_STYLE
    assert assistant._sanitize_chart_style({}) == assistant._DEFAULT_CHART_STYLE


def test_sanitize_chart_style_title_is_trimmed_and_length_capped():
    assert assistant._sanitize_chart_style({"title": "  Past 4 Months  "})["title"] == "Past 4 Months"
    assert assistant._sanitize_chart_style({"title": None})["title"] == ""
    long_title = "x" * 200
    assert len(assistant._sanitize_chart_style({"title": long_title})["title"]) == assistant._MAX_CHART_TITLE_LENGTH


def test_sanitize_chart_style_new_fields_valid_input():
    result = assistant._sanitize_chart_style({
        "show_legend": True,
        "y_axis_label": "  Revenue ($)  ",
        "opacity": 0.6,
        "corner_radius": 8,
        "font_size": "large",
        "show_data_points": True,
    })
    assert result["show_legend"] is True
    assert result["y_axis_label"] == "Revenue ($)"
    assert result["opacity"] == 0.6
    assert result["corner_radius"] == 8
    assert result["font_size"] == "large"
    assert result["show_data_points"] is True


def test_sanitize_chart_style_new_fields_adversarial_input():
    result = assistant._sanitize_chart_style({
        "show_legend": "yes",
        "y_axis_label": "x" * 200,
        "opacity": 5.0,  # way over the 1.0 max
        "corner_radius": -10,  # under the 0 min
        "font_size": "gigantic",
        "show_data_points": 1,
    })
    assert result["show_legend"] is True
    assert len(result["y_axis_label"]) == assistant._MAX_AXIS_LABEL_LENGTH
    assert result["opacity"] == 1.0
    assert result["corner_radius"] == 0
    assert result["font_size"] == "medium"
    assert result["show_data_points"] is True

    result2 = assistant._sanitize_chart_style({"opacity": "not-a-number", "corner_radius": "also-not-a-number"})
    assert result2["opacity"] == 1.0
    assert result2["corner_radius"] == 4


def test_resolve_chart_type_overrides_when_requested_and_different():
    chart = {"has_chart": True, "chart_type": "bar", "chart_categories": ["a"], "chart_series": []}
    result = assistant._resolve_chart_type(chart, {"chart_type": "line"})
    assert result["chart_type"] == "line"
    assert result["chart_categories"] == ["a"]  # untouched


def test_resolve_chart_type_leaves_auto_untouched():
    chart = {"has_chart": True, "chart_type": "bar", "chart_categories": ["a"], "chart_series": []}
    assert assistant._resolve_chart_type(chart, {"chart_type": "auto"}) == chart


def test_resolve_chart_type_does_nothing_without_a_chart():
    empty = assistant._empty_chart()
    assert assistant._resolve_chart_type(empty, {"chart_type": "line"}) == empty


# ---- chart_style: ask() round trip ------------------------------------------

def test_chart_style_defaults_when_model_omits_customization(monkeypatch):
    _install_fake_client(monkeypatch, [_final_answer()])

    result = assistant.ask("how did we do last week?", history=None, customer="fairprice")

    assert result["chart_style"] == assistant._DEFAULT_CHART_STYLE
    assert result["chart_trend_values"] == []


def test_chart_style_invalid_color_from_model_is_sanitized(monkeypatch):
    _install_fake_client(
        monkeypatch,
        [_final_answer(chart_style={"color": "mauve", "show_value_labels": False, "show_trend_line": False})],
    )

    result = assistant.ask("make it mauve", history=None, customer="fairprice")

    assert result["chart_style"]["color"] == "default"


def test_chart_style_valid_hex_color_passes_through(monkeypatch):
    _install_fake_client(
        monkeypatch,
        [_final_answer(chart_style={"color": "#DC2626", "show_value_labels": False, "show_trend_line": False})],
    )

    result = assistant.ask("make the bars red", history=None, customer="fairprice")

    assert result["chart_style"]["color"] == "#DC2626"


def test_trend_line_is_computed_from_the_charted_values(monkeypatch):
    summary = {
        "offline": {
            "storeFormats": [],
            "periodTotal": [
                {"period_label": "Week 1", "period_start": "2026-08-01", "revenue": 1000.0, "units": 10},
                {"period_label": "Week 2", "period_start": "2026-08-08", "revenue": 1500.0, "units": 15},
            ],
            "periodByFormat": [],
        },
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "week"}))
    _install_fake_client(
        monkeypatch,
        [tool_call, _final_answer(chart_style={"color": "default", "show_value_labels": False, "show_trend_line": True})],
    )

    result = assistant.ask("add a trend line", history=None, customer="fairprice")

    assert result["chart_style"]["show_trend_line"] is True
    assert result["chart_trend_values"] == [1000.0, 1500.0]


def test_no_trend_line_values_when_not_requested(monkeypatch):
    summary = {
        "offline": {
            "storeFormats": [],
            "periodTotal": [{"period_label": "Week 1", "period_start": "2026-08-01", "revenue": 1000.0, "units": 10}],
            "periodByFormat": [],
        },
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "week"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer()])

    result = assistant.ask("how did we do last week?", history=None, customer="fairprice")

    assert result["chart_trend_values"] == []


def test_last_chart_carried_forward_when_no_tool_call_this_turn(monkeypatch):
    # A pure styling follow-up ("make it red") makes no tool call -- the
    # chart the frontend already has on screen must still come back, not an
    # empty one, or there'd be nothing left to style.
    last_chart = {
        "has_chart": True, "chart_type": "bar", "chart_categories": ["May 2026"],
        "chart_series": [{"label": "Revenue", "values": [5473.15]}],
        "has_table": False, "table_columns": [], "table_rows": [],
    }
    _install_fake_client(
        monkeypatch,
        [_final_answer(
            uses_previous_chart=True,
            chart_style={"color": "#DC2626", "show_value_labels": False, "show_trend_line": False},
        )],
    )

    result = assistant.ask("make it red", history=None, customer="fairprice", last_chart=last_chart)

    assert result["has_chart"] is True
    assert result["chart_categories"] == ["May 2026"]
    assert result["chart_series"] == [{"label": "Revenue", "values": [5473.15]}]
    assert result["chart_style"]["color"] == "#DC2626"


def test_last_chart_not_carried_forward_for_an_off_topic_refusal(monkeypatch):
    # Regression (found via live testing): "what should I eat for lunch"
    # also makes no tool call, same as a styling follow-up -- but unlike
    # "make it red", it has nothing to do with the previous chart, so
    # carrying it forward showed a misleading graph under an answer that
    # was actually just a refusal.
    last_chart = {
        "has_chart": True, "chart_type": "bar", "chart_categories": ["May 2026"],
        "chart_series": [{"label": "Revenue", "values": [5473.15]}],
        "has_table": False, "table_columns": [], "table_rows": [],
    }
    _install_fake_client(
        monkeypatch,
        [_final_answer(
            text="I can't help with that -- I can only answer questions about AireOS's sales data.",
            grounded=True,
            data_source="none",
        )],
    )

    result = assistant.ask("what should i eat for lunch", history=None, customer="fairprice", last_chart=last_chart)

    assert result["has_chart"] is False
    assert result["chart_categories"] == []


def test_last_chart_style_seeds_the_system_prompt(monkeypatch):
    fake_client = _install_fake_client(monkeypatch, [_final_answer()])

    assistant.ask(
        "add a trend line",
        history=None,
        customer="fairprice",
        last_chart_style={"color": "#DC2626", "show_value_labels": False, "show_trend_line": False},
    )

    first_call = fake_client.models.calls[0]
    assert "color=#DC2626" in first_call["config"].system_instruction


# ---- chart_style: export builders -------------------------------------------

def test_export_pptx_reflects_custom_color_and_value_labels(monkeypatch):
    from pptx import Presentation
    import io as _io

    summary = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-01", "revenue": 5473.15, "units": 388}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month", "mode": "offline"}))
    _install_fake_client(
        monkeypatch,
        [
            tool_call,
            _final_answer(
                text="Offline sales in August 2026 were $5,473.15.",
                export_format="pptx",
                chart_style={"color": "#DC2626", "show_value_labels": True, "show_trend_line": False},
            ),
        ],
    )

    result = assistant.ask("export this as pptx in red with value labels", history=None, customer="fairprice")

    decoded = base64.b64decode(result["download_base64"])
    prs = Presentation(_io.BytesIO(decoded))
    chart_slide = prs.slides[1]
    chart = next(shape.chart for shape in chart_slide.shapes if shape.has_chart)
    assert str(chart.series[0].format.fill.fore_color.rgb) == "DC2626"
    assert chart.plots[0].has_data_labels is True


def test_export_pptx_reflects_value_format_and_gridlines(monkeypatch):
    from pptx import Presentation
    import io as _io

    summary = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-01", "revenue": 5473.15, "units": 388}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month", "mode": "offline"}))
    _install_fake_client(
        monkeypatch,
        [
            tool_call,
            _final_answer(
                export_format="pptx",
                chart_style={
                    **assistant._DEFAULT_CHART_STYLE,
                    "show_value_labels": True,
                    "value_label_format": "whole_number",
                    "show_gridlines": False,
                },
            ),
        ],
    )

    result = assistant.ask("export as pptx with whole numbers and no gridlines", history=None, customer="fairprice")

    decoded = base64.b64decode(result["download_base64"])
    prs = Presentation(_io.BytesIO(decoded))
    chart = next(shape.chart for shape in prs.slides[1].shapes if shape.has_chart)
    assert chart.plots[0].data_labels.number_format == "0"
    assert chart.value_axis.has_major_gridlines is False


def test_export_pptx_with_line_chart_type(monkeypatch):
    from pptx import Presentation
    from pptx.enum.chart import XL_CHART_TYPE
    import io as _io

    summary = {
        "offline": {
            "storeFormats": [],
            "periodTotal": [
                {"period_label": "Week 1", "period_start": "2026-08-01", "revenue": 1000.0, "units": 10},
                {"period_label": "Week 2", "period_start": "2026-08-08", "revenue": 1500.0, "units": 15},
            ],
            "periodByFormat": [],
        },
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "week"}))
    _install_fake_client(
        monkeypatch,
        [tool_call, _final_answer(export_format="pptx", chart_style={**assistant._DEFAULT_CHART_STYLE, "chart_type": "line", "color": "#DC2626"})],
    )

    result = assistant.ask("show this as a line chart and export as pptx", history=None, customer="fairprice")

    assert result["chart_type"] == "line"
    decoded = base64.b64decode(result["download_base64"])
    prs = Presentation(_io.BytesIO(decoded))
    chart = next(shape.chart for shape in prs.slides[1].shapes if shape.has_chart)
    assert chart.chart_type == XL_CHART_TYPE.LINE
    assert str(chart.series[0].format.line.color.rgb) == "DC2626"


def test_chart_data_is_never_affected_by_adversarial_chart_style(monkeypatch):
    # The concrete guarantee behind "strictly from BigQuery": however wild
    # chart_style gets, chart_categories/chart_series must come out exactly
    # as _build_chart produced them -- proven here across a battery of
    # malformed/adversarial style payloads, not just the happy path.
    summary = {
        "offline": {
            "storeFormats": [],
            "periodTotal": [
                {"period_label": "Week 1", "period_start": "2026-08-01", "revenue": 1000.0, "units": 10},
                {"period_label": "Week 2", "period_start": "2026-08-08", "revenue": 1500.0, "units": 15},
            ],
            "periodByFormat": [],
        },
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    adversarial_styles = [
        {"color": "javascript:alert(1)", "chart_type": "pie", "value_label_format": "$$$hack$$$"},
        {"color": "#DC2626; DROP TABLE sales;", "show_value_labels": "definitely", "chart_type": "scatter3d"},
        {"chart_type": "line", "value_label_format": "whole_number", "show_gridlines": "not-a-bool"},
        {"title": "x" * 500},
        {"opacity": -99, "corner_radius": "banana", "font_size": "gigantic", "y_axis_label": "x" * 999, "show_legend": "maybe"},
        {},  # nothing at all
    ]

    for style in adversarial_styles:
        tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "week"}))
        _install_fake_client(monkeypatch, [tool_call, _final_answer(chart_style=style)])

        result = assistant.ask("how is revenue trending?", history=None, customer="fairprice")

        assert result["chart_categories"] == ["Week 1", "Week 2"]
        assert result["chart_series"] == [{"label": "Revenue", "values": [1000.0, 1500.0]}]


def test_export_pdf_with_custom_color_still_produces_valid_pdf(monkeypatch):
    summary = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-01", "revenue": 5473.15, "units": 388}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month", "mode": "offline"}))
    _install_fake_client(
        monkeypatch,
        [
            tool_call,
            _final_answer(
                text="Offline sales in August 2026 were $5,473.15.",
                export_format="pdf",
                chart_style={"color": "#DC2626", "show_value_labels": True, "show_trend_line": False},
            ),
        ],
    )

    result = assistant.ask("export this as pdf in red", history=None, customer="fairprice")

    decoded = base64.b64decode(result["download_base64"])
    assert decoded[:4] == b"%PDF"


def test_export_pptx_reflects_custom_chart_title(monkeypatch):
    from pptx import Presentation
    import io as _io

    summary = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-01", "revenue": 5473.15, "units": 388}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month", "mode": "offline"}))
    _install_fake_client(
        monkeypatch,
        [tool_call, _final_answer(export_format="pptx", chart_style={**assistant._DEFAULT_CHART_STYLE, "title": "Past 4 Months"})],
    )

    result = assistant.ask("add a title called Past 4 Months and export as pptx", history=None, customer="fairprice")

    assert result["chart_style"]["title"] == "Past 4 Months"
    decoded = base64.b64decode(result["download_base64"])
    prs = Presentation(_io.BytesIO(decoded))
    chart = next(shape.chart for shape in prs.slides[1].shapes if shape.has_chart)
    assert chart.has_title is True
    assert chart.chart_title.text_frame.text == "Past 4 Months"


def test_export_pdf_with_custom_title_still_produces_valid_pdf(monkeypatch):
    summary = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-01", "revenue": 5473.15, "units": 388}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month", "mode": "offline"}))
    _install_fake_client(
        monkeypatch,
        [tool_call, _final_answer(export_format="pdf", chart_style={**assistant._DEFAULT_CHART_STYLE, "title": "Past 4 Months"})],
    )

    result = assistant.ask("add a title called Past 4 Months and export as pdf", history=None, customer="fairprice")

    decoded = base64.b64decode(result["download_base64"])
    assert decoded[:4] == b"%PDF"


def test_export_pptx_reflects_legend_axis_title_currency_and_markers(monkeypatch):
    from pptx import Presentation
    from pptx.enum.chart import XL_MARKER_STYLE
    import io as _io

    summary = {
        "offline": {
            "storeFormats": [],
            "periodTotal": [
                {"period_label": "Week 1", "period_start": "2026-08-01", "revenue": 1000.0, "units": 10},
                {"period_label": "Week 2", "period_start": "2026-08-08", "revenue": 1500.0, "units": 15},
            ],
            "periodByFormat": [],
        },
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "week"}))
    style = {
        **assistant._DEFAULT_CHART_STYLE,
        "chart_type": "line",
        "show_legend": True,
        "y_axis_label": "Revenue ($)",
        "show_value_labels": True,
        "value_label_format": "currency",
        "show_data_points": True,
    }
    _install_fake_client(monkeypatch, [tool_call, _final_answer(export_format="pptx", chart_style=style)])

    result = assistant.ask("export with legend, axis label, currency labels and markers", history=None, customer="fairprice")

    decoded = base64.b64decode(result["download_base64"])
    prs = Presentation(_io.BytesIO(decoded))
    chart = next(shape.chart for shape in prs.slides[1].shapes if shape.has_chart)
    assert chart.has_legend is True
    assert chart.value_axis.has_title is True
    assert chart.value_axis.axis_title.text_frame.text == "Revenue ($)"
    assert chart.plots[0].data_labels.number_format == '"$"#,##0'
    assert chart.series[0].marker.style == XL_MARKER_STYLE.CIRCLE


def test_export_pdf_with_opacity_currency_and_markers_still_produces_valid_pdf(monkeypatch):
    summary = {
        "offline": {
            "storeFormats": [],
            "periodTotal": [
                {"period_label": "Week 1", "period_start": "2026-08-01", "revenue": 1000.0, "units": 10},
                {"period_label": "Week 2", "period_start": "2026-08-08", "revenue": 1500.0, "units": 15},
            ],
            "periodByFormat": [],
        },
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "week"}))
    style = {
        **assistant._DEFAULT_CHART_STYLE,
        "chart_type": "line",
        "opacity": 0.5,
        "y_axis_label": "Revenue ($)",
        "show_value_labels": True,
        "value_label_format": "currency",
        "show_data_points": True,
        "font_size": "large",
    }
    _install_fake_client(monkeypatch, [tool_call, _final_answer(export_format="pdf", chart_style=style)])

    result = assistant.ask("export with opacity, axis label, currency labels and markers", history=None, customer="fairprice")

    decoded = base64.b64decode(result["download_base64"])
    assert decoded[:4] == b"%PDF"


def test_export_pptx_bar_chart_includes_trend_line_as_combo_chart(monkeypatch):
    # Regression: trend line previously only showed in the chat preview,
    # not in exports at all. Verifies the real OOXML combo-chart structure
    # (a second <c:lineChart> sibling under <c:plotArea>, not just "some
    # bytes got written") round-trips correctly.
    from pptx import Presentation
    from pptx.oxml.ns import qn
    import io as _io

    summary = {
        "offline": {
            "storeFormats": [],
            "periodTotal": [
                {"period_label": "Week 1", "period_start": "2026-08-01", "revenue": 1000.0, "units": 10},
                {"period_label": "Week 2", "period_start": "2026-08-08", "revenue": 1500.0, "units": 15},
                {"period_label": "Week 3", "period_start": "2026-08-15", "revenue": 1200.0, "units": 12},
                {"period_label": "Week 4", "period_start": "2026-08-22", "revenue": 1800.0, "units": 18},
            ],
            "periodByFormat": [],
        },
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "week"}))
    style = {**assistant._DEFAULT_CHART_STYLE, "show_trend_line": True}
    _install_fake_client(monkeypatch, [tool_call, _final_answer(export_format="pptx", chart_style=style)])

    result = assistant.ask("add a trend line and export as pptx", history=None, customer="fairprice")

    assert result["chart_trend_values"], "trend values should have been computed"
    decoded = base64.b64decode(result["download_base64"])
    prs = Presentation(_io.BytesIO(decoded))
    chart = next(shape.chart for shape in prs.slides[1].shapes if shape.has_chart)
    plot_area = chart._chartSpace.chart.plotArea
    assert plot_area.find(qn("c:barChart")) is not None
    line_elm = plot_area.find(qn("c:lineChart"))
    assert line_elm is not None
    trend_vals = [pt.find(qn("c:v")).text for pt in line_elm.findall(".//" + qn("c:val") + "/" + qn("c:numRef") + "/" + qn("c:numCache") + "/" + qn("c:pt"))]
    assert [float(v) for v in trend_vals] == result["chart_trend_values"]
    assert chart.series[0].values == (1000.0, 1500.0, 1200.0, 1800.0)  # primary series untouched


def test_export_pptx_line_chart_includes_second_trend_series(monkeypatch):
    from pptx import Presentation
    from pptx.oxml.ns import qn
    import io as _io

    summary = {
        "offline": {
            "storeFormats": [],
            "periodTotal": [
                {"period_label": "Week 1", "period_start": "2026-08-01", "revenue": 1000.0, "units": 10},
                {"period_label": "Week 2", "period_start": "2026-08-08", "revenue": 1500.0, "units": 15},
                {"period_label": "Week 3", "period_start": "2026-08-15", "revenue": 1200.0, "units": 12},
            ],
            "periodByFormat": [],
        },
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "week"}))
    style = {**assistant._DEFAULT_CHART_STYLE, "show_trend_line": True, "chart_type": "line"}
    _install_fake_client(monkeypatch, [tool_call, _final_answer(export_format="pptx", chart_style=style)])

    result = assistant.ask("show as line chart with trend line, export pptx", history=None, customer="fairprice")

    decoded = base64.b64decode(result["download_base64"])
    prs = Presentation(_io.BytesIO(decoded))
    chart = next(shape.chart for shape in prs.slides[1].shapes if shape.has_chart)
    plot_area = chart._chartSpace.chart.plotArea
    line_charts = plot_area.findall(qn("c:lineChart"))
    assert len(line_charts) == 1  # one lineChart element containing TWO series, not two lineChart siblings
    assert len(line_charts[0].findall(qn("c:ser"))) == 2
    assert chart.series[0].values == (1000.0, 1500.0, 1200.0)


def test_export_pdf_bar_chart_with_trend_line_still_produces_valid_pdf(monkeypatch):
    summary = {
        "offline": {
            "storeFormats": [],
            "periodTotal": [
                {"period_label": "Week 1", "period_start": "2026-08-01", "revenue": 1000.0, "units": 10},
                {"period_label": "Week 2", "period_start": "2026-08-08", "revenue": 1500.0, "units": 15},
            ],
            "periodByFormat": [],
        },
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "week"}))
    style = {**assistant._DEFAULT_CHART_STYLE, "show_trend_line": True}
    _install_fake_client(monkeypatch, [tool_call, _final_answer(export_format="pdf", chart_style=style)])

    result = assistant.ask("add a trend line and export as pdf", history=None, customer="fairprice")

    decoded = base64.b64decode(result["download_base64"])
    assert decoded[:4] == b"%PDF"


def test_export_pdf_line_chart_with_trend_line_still_produces_valid_pdf(monkeypatch):
    summary = {
        "offline": {
            "storeFormats": [],
            "periodTotal": [
                {"period_label": "Week 1", "period_start": "2026-08-01", "revenue": 1000.0, "units": 10},
                {"period_label": "Week 2", "period_start": "2026-08-08", "revenue": 1500.0, "units": 15},
            ],
            "periodByFormat": [],
        },
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "week"}))
    style = {**assistant._DEFAULT_CHART_STYLE, "show_trend_line": True, "chart_type": "line"}
    _install_fake_client(monkeypatch, [tool_call, _final_answer(export_format="pdf", chart_style=style)])

    result = assistant.ask("show as line with trend line, export as pdf", history=None, customer="fairprice")

    decoded = base64.b64decode(result["download_base64"])
    assert decoded[:4] == b"%PDF"


def test_export_pdf_with_line_chart_type_still_produces_valid_pdf(monkeypatch):
    summary = {
        "offline": {
            "storeFormats": [],
            "periodTotal": [
                {"period_label": "Week 1", "period_start": "2026-08-01", "revenue": 1000.0, "units": 10},
                {"period_label": "Week 2", "period_start": "2026-08-08", "revenue": 1500.0, "units": 15},
            ],
            "periodByFormat": [],
        },
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "week"}))
    _install_fake_client(
        monkeypatch,
        [tool_call, _final_answer(export_format="pdf", chart_style={**assistant._DEFAULT_CHART_STYLE, "chart_type": "line"})],
    )

    result = assistant.ask("show as a line chart and export as pdf", history=None, customer="fairprice")

    assert result["chart_type"] == "line"
    decoded = base64.b64decode(result["download_base64"])
    assert decoded[:4] == b"%PDF"


# ---- get_promotions: filtering helpers --------------------------------------
# A promotion can run at more than one store (a real many-to-many via
# promotion_stores, confirmed against live data -- 3 of 5 real promotions in
# the dev DB link to 2 stores each), so `stores` is a list per promotion.

def _sample_promotion(**overrides):
    promo = {
        "promotion_id": 1,
        "period_start": datetime.date(2026, 4, 1),
        "period_end": datetime.date(2026, 4, 30),
        "period_label": "Apr 2026",
        "promo_type": "regular",
        "promotion_mechanic": "20% off",
        "voucher": None,
        "stores": [
            {"store_id": 1, "store_name": "FairPrice Jurong Point", "store_code": "JP01", "store_format": "Hypermarket", "retailer_id": 1, "retailer": "FairPrice"},
        ],
        "skus": [
            {"sku": "AP-SM-001", "sku_range": "Aire Adult Pants", "product_name": "Aire Adult Pants S/M", "size": "S/M", "brand": "Aire", "uom": "pack", "pack_size": 10, "price": 12.90, "quantity_units": None},
        ],
        "created_at": datetime.datetime(2026, 3, 1, 9, 0, 0),
        "updated_at": datetime.datetime(2026, 3, 1, 9, 0, 0),
    }
    promo.update(overrides)
    return promo


def test_promotion_matches_customer_checks_any_linked_store():
    promo = _sample_promotion(stores=[
        {"retailer": "FairPrice", "store_name": "Jurong Point"},
        {"retailer": "ColdStorage", "store_name": "Orchard"},
    ])
    assert assistant._promotion_matches_customer(promo, "fairprice") is True
    assert assistant._promotion_matches_customer(promo, "COLD") is True
    assert assistant._promotion_matches_customer(promo, "sheng siong") is False
    assert assistant._promotion_matches_customer(promo, None) is True


def test_promotion_overlaps_range():
    promo = _sample_promotion(period_start=datetime.date(2026, 4, 1), period_end=datetime.date(2026, 4, 30))
    assert assistant._promotion_overlaps_range(promo, "2026-04-15", "2026-04-20") is True  # fully inside
    assert assistant._promotion_overlaps_range(promo, "2026-03-01", "2026-04-05") is True  # overlaps start
    assert assistant._promotion_overlaps_range(promo, "2026-04-25", "2026-05-10") is True  # overlaps end
    assert assistant._promotion_overlaps_range(promo, "2026-05-01", "2026-05-31") is False  # entirely after
    assert assistant._promotion_overlaps_range(promo, "2026-01-01", "2026-01-31") is False  # entirely before
    assert assistant._promotion_overlaps_range(promo, None, None) is True


def test_promotion_matches_sku_checks_code_and_product_name():
    promo = _sample_promotion(skus=[{"sku": "AP-SM-001", "product_name": "Aire Adult Pants S/M"}])
    assert assistant._promotion_matches_sku(promo, "AP-SM-001") is True
    assert assistant._promotion_matches_sku(promo, "adult pants") is True
    assert assistant._promotion_matches_sku(promo, "widget") is False
    assert assistant._promotion_matches_sku(promo, None) is True


def test_trim_promotion_drops_dashboard_detail_noise_and_keeps_all_stores():
    trimmed = assistant._trim_promotion(_sample_promotion(stores=[
        {"store_id": 1, "store_code": "JP01", "store_format": "Hypermarket", "retailer_id": 1, "retailer": "FairPrice", "store_name": "Jurong Point"},
        {"store_id": 2, "store_code": "OR02", "store_format": "Supermarket", "retailer_id": 1, "retailer": "FairPrice", "store_name": "Orchard"},
    ]))
    assert "price" not in json.dumps(trimmed, default=str)
    assert "created_at" not in trimmed
    assert "updated_at" not in trimmed
    assert trimmed["stores"] == [
        {"retailer": "FairPrice", "store_name": "Jurong Point"},
        {"retailer": "FairPrice", "store_name": "Orchard"},
    ]
    assert "store_id" not in trimmed["stores"][0]
    assert trimmed["skus"] == ["Aire Adult Pants Aire Adult Pants S/M"]
    assert trimmed["promotion_mechanic"] == "20% off"


# ---- get_promotions: ask() integration --------------------------------------

def test_get_promotions_tool_filters_by_customer_and_date_range(monkeypatch):
    promotions = [
        _sample_promotion(period_label="Apr 2026", period_start=datetime.date(2026, 4, 1), period_end=datetime.date(2026, 4, 30), stores=[{"retailer": "FairPrice", "store_name": "Jurong Point"}]),
        _sample_promotion(period_label="Apr 2026", period_start=datetime.date(2026, 4, 1), period_end=datetime.date(2026, 4, 30), stores=[{"retailer": "ColdStorage", "store_name": "Orchard"}]),
        _sample_promotion(period_label="Jan 2026", period_start=datetime.date(2026, 1, 1), period_end=datetime.date(2026, 1, 31), stores=[{"retailer": "FairPrice", "store_name": "Jurong Point"}]),
    ]
    monkeypatch.setattr(promotion_service, "get_promotions", lambda: promotions)

    tool_call = _tool_call_response(("tu_1", "get_promotions", {"customer": "fairprice", "start_date": "2026-04-01", "end_date": "2026-04-30"}))
    fake_client = _install_fake_client(monkeypatch, [tool_call, _final_answer(text="FairPrice ran a 20% off promotion in April 2026.")])

    result = assistant.ask("what promotions did fairprice run in april 2026?", history=None, customer="fairprice")

    tool_result_message = _dump(fake_client.models.calls[1]["contents"][-1])
    function_response = json.loads(tool_result_message["parts"][0]["function_response"]["response"]["result"])
    assert len(function_response) == 1
    assert function_response[0]["stores"][0]["retailer"] == "FairPrice"
    assert function_response[0]["period_label"] == "Apr 2026"
    assert result["has_table"] is True
    assert result["table_columns"] == ["stores", "period", "type", "mechanic"]


def test_get_promotions_tool_empty_result_still_answers(monkeypatch):
    monkeypatch.setattr(promotion_service, "get_promotions", lambda: [])

    tool_call = _tool_call_response(("tu_1", "get_promotions", {"customer": "fairprice"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="No promotions found for FairPrice.")])

    result = assistant.ask("what promotions has fairprice run?", history=None, customer="fairprice")

    assert result["answer"] == "No promotions found for FairPrice."
    assert result["has_table"] is False


def test_get_promotions_tool_db_failure_degrades_gracefully(monkeypatch):
    def _raise():
        raise RuntimeError("connection refused")
    monkeypatch.setattr(promotion_service, "get_promotions", _raise)

    tool_call = _tool_call_response(("tu_1", "get_promotions", {"customer": "fairprice"}))
    fake_client = _install_fake_client(
        monkeypatch,
        [tool_call, _final_answer(grounded=True, data_source="none", text="I couldn't reach the promotions database.")],
    )

    result = assistant.ask("what promotions has fairprice run?", history=None, customer="fairprice")

    tool_result_message = _dump(fake_client.models.calls[1]["contents"][-1])
    function_response = tool_result_message["parts"][0]["function_response"]
    assert "error" in function_response["response"]
    assert result["answer"] == "I couldn't reach the promotions database."


def test_get_promotions_builds_table_with_multiple_stores_joined(monkeypatch):
    promotions = [_sample_promotion(
        promo_type="side_offer",
        promotion_mechanic="Buy 1 Get 1",
        stores=[
            {"retailer": "FairPrice", "store_name": "Jurong Point"},
            {"retailer": "FairPrice", "store_name": "Orchard"},
        ],
    )]
    monkeypatch.setattr(promotion_service, "get_promotions", lambda: promotions)

    tool_call = _tool_call_response(("tu_1", "get_promotions", {"customer": "fairprice"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer()])

    result = assistant.ask("what promotions has fairprice run?", history=None, customer="fairprice")

    assert result["has_chart"] is False
    assert result["has_table"] is True
    assert result["table_rows"][0] == ["FairPrice - Jurong Point; FairPrice - Orchard", "Apr 2026", "side_offer", "Buy 1 Get 1"]


def test_get_promotions_tool_scoped_by_sku(monkeypatch):
    promotions = [
        _sample_promotion(skus=[{"sku": "AP-SM-001", "product_name": "Aire Adult Pants S/M"}]),
        _sample_promotion(skus=[{"sku": "AP-XL-001", "product_name": "Aire Adult Pants XL"}]),
    ]
    monkeypatch.setattr(promotion_service, "get_promotions", lambda: promotions)

    tool_call = _tool_call_response(("tu_1", "get_promotions", {"customer": "fairprice", "sku": "XL"}))
    fake_client = _install_fake_client(monkeypatch, [tool_call, _final_answer()])

    assistant.ask("was the XL pants sku ever on promotion?", history=None, customer="fairprice")

    tool_result_message = _dump(fake_client.models.calls[1]["contents"][-1])
    function_response = json.loads(tool_result_message["parts"][0]["function_response"]["response"]["result"])
    assert len(function_response) == 1
    assert "XL" in function_response[0]["skus"][0]


# ---- generate_digest: pure helpers -------------------------------------------

def test_rank_movers_picks_top_n_by_absolute_delta():
    current = [
        {"sku": "A", "product_name": "Widget A", "rank": 1},
        {"sku": "B", "product_name": "Widget B", "rank": 2},
        {"sku": "C", "product_name": "Widget C", "rank": 3},
        {"sku": "D", "product_name": "Widget D", "rank": 4},
    ]
    previous = [
        {"sku": "A", "rank": 1},   # no change -- excluded
        {"sku": "B", "rank": 5},   # moved up 3
        {"sku": "C", "rank": 2},   # moved down 1
        {"sku": "D", "rank": 10},  # moved up 6 -- biggest mover
    ]

    movers = assistant._rank_movers(current, previous, limit=2)

    assert [m["sku"] for m in movers] == ["D", "B"]
    assert movers[0]["delta"] == 6
    assert movers[0]["previous_rank"] == 10
    assert movers[0]["current_rank"] == 4


def test_rank_movers_skips_skus_missing_from_either_period():
    current = [{"sku": "A", "product_name": "Widget A", "rank": 1}]
    previous = [{"sku": "B", "rank": 1}]  # different SKU entirely

    assert assistant._rank_movers(current, previous) == []


def test_promotions_starting_or_ending_soon_windows_by_date_and_customer():
    today = datetime.date.today()
    promos = [
        _sample_promotion(  # starts in 3 days -- within window
            period_start=today + datetime.timedelta(days=3),
            period_end=today + datetime.timedelta(days=30),
            stores=[{"retailer": "FairPrice", "store_name": "Jurong Point"}],
        ),
        _sample_promotion(  # ends in 5 days -- within window
            period_start=today - datetime.timedelta(days=20),
            period_end=today + datetime.timedelta(days=5),
            stores=[{"retailer": "FairPrice", "store_name": "Orchard"}],
        ),
        _sample_promotion(  # far in the future -- outside the 7-day window
            period_start=today + datetime.timedelta(days=60),
            period_end=today + datetime.timedelta(days=90),
            stores=[{"retailer": "FairPrice", "store_name": "Bedok"}],
        ),
        _sample_promotion(  # in window but wrong retailer
            period_start=today + datetime.timedelta(days=1),
            period_end=today + datetime.timedelta(days=10),
            stores=[{"retailer": "ColdStorage", "store_name": "Somewhere"}],
        ),
    ]

    result = assistant._promotions_starting_or_ending_soon(promos, "fairprice", days=7)

    stores_seen = {s["store_name"] for p in result for s in p["stores"]}
    assert stores_seen == {"Jurong Point", "Orchard"}


# ---- generate_digest: the month window ------------------------------------------

def test_digest_window_compares_a_full_month_with_the_whole_month_before():
    window = assistant._digest_window(datetime.date(2026, 8, 1), datetime.date(2026, 8, 31))

    assert (window["current_start"], window["current_end"]) == ("2026-08-01", "2026-08-31")
    assert (window["previous_start"], window["previous_end"]) == ("2026-07-01", "2026-07-31")
    assert (window["previous_label"], window["current_label"]) == ("Jul 2026", "Aug 2026")
    assert window["full_month"] is True


def test_digest_window_compares_a_part_month_with_the_same_days_of_the_month_before():
    # Weekly data up to 15 Sep: 1-15 Sep against 1-15 Aug, not against all of August.
    window = assistant._digest_window(datetime.date(2026, 9, 9), datetime.date(2026, 9, 15))

    assert (window["current_start"], window["current_end"]) == ("2026-09-01", "2026-09-15")
    assert (window["previous_start"], window["previous_end"]) == ("2026-08-01", "2026-08-15")
    assert window["current_label"] == "1-15 Sep 2026 (month to date)"
    assert window["full_month"] is False


def test_digest_window_never_runs_past_the_end_of_a_shorter_previous_month():
    window = assistant._digest_window(datetime.date(2026, 3, 1), datetime.date(2026, 3, 31))

    assert (window["previous_start"], window["previous_end"]) == ("2026-02-01", "2026-02-28")


def test_digest_window_uses_the_month_a_week_starts_in():
    # A week from 30 Jul to 5 Aug counts towards July, as every sales read does.
    window = assistant._digest_window(datetime.date(2026, 7, 30), datetime.date(2026, 8, 5))

    assert (window["current_start"], window["current_end"]) == ("2026-07-01", "2026-07-31")
    assert window["full_month"] is True


def test_next_month_rolls_over_the_year():
    assert assistant._next_month(datetime.date(2026, 10, 29)) == "2026-11"
    assert assistant._next_month(datetime.date(2026, 12, 31)) == "2027-01"


def test_store_summary_names_a_few_stores_and_counts_many():
    few = [{"retailer": "fairprice_offline", "store_name": "Bedok"}, {"retailer": "fairprice_online", "store_name": "Web"}]
    many = [{"retailer": "fairprice_offline", "store_name": f"Store {i}"} for i in range(53)]
    many.append({"retailer": "fairprice_online", "store_name": "Web"})

    assert assistant._store_summary(few) == "fairprice_offline - Bedok, fairprice_online - Web"
    assert assistant._store_summary(many) == "53 fairprice_offline stores, 1 fairprice_online store"


def test_change_words_says_up_down_or_unchanged_instead_of_signs():
    assert assistant._change_words(120.0, 100.0) == "up 20.0%"
    assert assistant._change_words(88.8, 100.0) == "down 11.2%"
    assert assistant._change_words(100.01, 100.0) == "unchanged"
    assert assistant._change_words(50.0, 0.0) is None


def test_promotion_lines_give_a_count_and_readable_dates():
    lines = assistant._digest_promotion_lines([{
        "promotion_mechanic": "20% Off",
        "promo_type": "monthly",
        "stores": [{"retailer": "fairprice_offline", "store_name": "Bedok"}],
        "period_start": datetime.date(2026, 11, 1),
        "period_end": datetime.date(2026, 11, 30),
    }])

    assert lines == [
        "- 1 promotion starting or ending in the next 7 days:",
        "  - 20% Off (monthly) at fairprice_offline - Bedok, running 1 Nov 2026 to 30 Nov 2026.",
    ]


# ---- generate_digest: end-to-end ---------------------------------------------

def _digest_answer(text="- Revenue trend was strong.", follow_ups=None):
    payload = {
        "answer": text,
        "follow_up_prompts": follow_ups or ["What drove this?", "Show me last month too?"],
    }
    return FakeResponse(
        text=json.dumps(payload),
        content={"role": "model", "parts": [{"text": json.dumps(payload)}]},
    )


def _comparison(current_revenue, previous_revenue, current_units=0.0, previous_units=0.0, available=True):
    return {
        "current": {"start": "2026-08-01", "end": "2026-08-31", "revenue": current_revenue, "units": current_units},
        "previous": {
            "start": "2026-07-01", "end": "2026-07-31",
            "revenue": previous_revenue, "units": previous_units, "available": available,
        },
    }


def _install_digest_data(
    monkeypatch,
    latest_by_channel=None,
    comparisons=None,
    rankings=None,
    at_risk=None,
    forecast=None,
    promotions=None,
):
    """Every read generate_digest makes, faked. Returns the period comparison calls it saw."""
    if latest_by_channel is None:
        latest_by_channel = {
            "offline": ("2026-08-01", "2026-08-31"),
            "online": ("2026-07-30", "2026-08-05"),
        }
    comparisons = comparisons or {
        None: _comparison(12000.0, 10000.0, 900.0, 800.0),
        "offline": _comparison(11000.0, 9000.0),
        "online": _comparison(1000.0, 1000.0),
    }
    comparison_calls = []

    def fake_default_range(customer, mode, period):
        start, end = latest_by_channel.get(mode) or (None, None)
        return {"start": start, "end": end}

    def fake_comparison(**kwargs):
        comparison_calls.append(kwargs)
        return comparisons[kwargs.get("mode")]

    def fake_ranking(**kwargs):
        rows = (rankings or {}).get(kwargs.get("start_date"), [])
        return pd.DataFrame(rows)

    monkeypatch.setattr(sellout_service, "get_default_date_range", fake_default_range)
    monkeypatch.setattr(sellout_service, "get_period_comparison", fake_comparison)
    monkeypatch.setattr(sellout_service, "get_sku_ranking", fake_ranking)
    monkeypatch.setattr(inventory_service, "list_customers", lambda: [{"customer_id": 1, "customer_name": "fairprice"}])
    monkeypatch.setattr(
        inventory_service, "get_at_risk",
        lambda customer_ids: at_risk or {"as_of": None, "counts": {}, "items": []},
    )

    def fake_forecast(customer, start_month):
        if forecast is None:
            raise ValueError("There is no forecast for that product and customer yet.")
        return forecast

    monkeypatch.setattr(forecast_summary, "get_demand_forecast", fake_forecast)
    monkeypatch.setattr(promotion_service, "get_promotions", lambda: promotions or [])
    return comparison_calls


def _facts_sent(fake_client) -> str:
    return fake_client.models.calls[0]["contents"][0]["parts"][0]["text"]


def test_generate_digest_compares_the_latest_month_with_the_month_before(monkeypatch):
    # Offline has all of August as a monthly total, online only weekly rows to
    # 5 Aug: the newest data sets the month, so it's August vs July.
    calls = _install_digest_data(monkeypatch)
    fake_client = _install_fake_client(monkeypatch, [_digest_answer()])

    result = assistant.generate_digest(customer="fairprice")

    assert {(c["current_start"], c["current_end"], c["previous_start"], c["previous_end"]) for c in calls} == {
        ("2026-08-01", "2026-08-31", "2026-07-01", "2026-07-31")
    }
    facts = _facts_sent(fake_client)
    assert "Sales in Aug 2026: revenue was $12,000, up 20.0% from $10,000 in Jul 2026." in facts
    assert "Units sold were 900, up 12.5% from 800 in Jul 2026." in facts
    assert "In-store (offline) $11,000, up 22.2%; Online $1,000, unchanged." in facts
    assert result["chart_categories"] == ["Jul 2026", "Aug 2026"]
    assert result["chart_series"] == [{"label": "Revenue", "values": [10000.0, 12000.0]}]
    assert result["grounded"] is True
    assert result["answer"] == "- Revenue trend was strong."


def test_generate_digest_names_the_best_seller_and_rank_movers(monkeypatch):
    _install_digest_data(monkeypatch, rankings={
        "2026-08-01": [
            {"sku": "A", "product_name": "Widget A", "rank": 1, "value": 5000.0},
            {"sku": "B", "product_name": "Widget B", "rank": 2, "value": 3000.0},
        ],
        "2026-07-01": [
            {"sku": "A", "product_name": "Widget A", "rank": 5, "value": 900.0},
            {"sku": "B", "product_name": "Widget B", "rank": 2, "value": 3000.0},
        ],
    })
    fake_client = _install_fake_client(monkeypatch, [_digest_answer()])

    assistant.generate_digest(customer="fairprice")

    facts = _facts_sent(fake_client)
    assert "Best-selling product in Aug 2026: Widget A, $5,000 in revenue." in facts
    assert "Widget A rose from #5 to #1." in facts
    assert "Widget B" not in facts.split("sales ranking")[1]


def test_generate_digest_includes_stock_outside_its_range(monkeypatch):
    _install_digest_data(monkeypatch, at_risk={
        "as_of": "2026-09-01",
        "counts": {"below_min": 2, "above_max": 1},
        "items": [
            {"sku": "A", "product_name": "Widget A", "doh": 5.0, "doh_status": "below_min", "days_outside_band": 20.0},
            {"sku": "B", "product_name": "Widget B", "doh": 50.0, "doh_status": "above_max", "days_outside_band": 15.0},
            {"sku": "C", "product_name": "Widget C", "doh": 20.0, "doh_status": "below_min", "days_outside_band": 5.0},
            {"sku": "D", "product_name": "Widget D", "doh": 21.0, "doh_status": "below_min", "days_outside_band": 4.0},
        ],
    })
    fake_client = _install_fake_client(monkeypatch, [_digest_answer()])

    assistant.generate_digest(customer="fairprice")

    facts = _facts_sent(fake_client)
    assert "Stock at the end of Sep 2026: 2 products running low (may run out)" in facts
    assert "1 product with more stock than needed" in facts
    assert "Widget A: 5 days of stock, 20 days short of the minimum." in facts
    assert "Widget B: 50 days of stock, 15 days over the maximum." in facts
    assert "Widget D" not in facts  # only the three furthest out


def test_generate_digest_includes_the_next_forecast_month_against_the_last_full_actual_month(monkeypatch):
    _install_digest_data(monkeypatch, forecast={
        "series": [{}, {}],
        "monthly_totals": [
            {"month": "2026-08-01", "forecast_cartons": None, "actual_cartons": 2400, "actual_month_to_date": False},
            {"month": "2026-09-01", "forecast_cartons": None, "actual_cartons": 300, "actual_month_to_date": True},
            {"month": "2026-11-01", "forecast_cartons": 2582, "actual_cartons": None, "actual_month_to_date": False},
            {"month": "2026-12-01", "forecast_cartons": 2700, "actual_cartons": None, "actual_month_to_date": False},
        ],
    })
    fake_client = _install_fake_client(monkeypatch, [_digest_answer()])

    assistant.generate_digest(customer="fairprice")

    assert (
        "Forecast for Nov 2026: about 2,582 cartons expected to sell, "
        "up 7.6% on the 2,400 cartons actually sold in Aug 2026."
    ) in _facts_sent(fake_client)


def test_generate_digest_includes_promotions_starting_soon(monkeypatch):
    today = datetime.date.today()
    _install_digest_data(monkeypatch, promotions=[_sample_promotion(
        period_start=today + datetime.timedelta(days=1),
        period_end=today + datetime.timedelta(days=10),
        promotion_mechanic="20% off",
        stores=[{"retailer": "FairPrice", "store_name": "Jurong Point"}],
    )])
    fake_client = _install_fake_client(monkeypatch, [_digest_answer()])

    assistant.generate_digest(customer="fairprice")

    assert "20% off" in _facts_sent(fake_client)


def test_generate_digest_leaves_out_an_area_it_cannot_read(monkeypatch):
    _install_digest_data(monkeypatch)

    def _down(**kwargs):
        raise RuntimeError("connection refused")

    monkeypatch.setattr(inventory_service, "get_at_risk", _down)
    monkeypatch.setattr(promotion_service, "get_promotions", lambda: _down())
    fake_client = _install_fake_client(monkeypatch, [_digest_answer()])

    result = assistant.generate_digest(customer="fairprice")

    facts = _facts_sent(fake_client)
    assert "Stock" not in facts
    assert "romotion" not in facts
    assert "connection refused" not in facts
    assert "Sales in Aug 2026" in facts  # the rest is still there
    assert result["has_chart"] is True


def test_generate_digest_turns_a_literal_backslash_n_into_a_line_break(monkeypatch):
    _install_digest_data(monkeypatch)
    _install_fake_client(monkeypatch, [_digest_answer(text="- **Sales:** fell.\\n- **Stock:** fine.")])

    result = assistant.generate_digest(customer="fairprice")

    assert result["answer"] == "- **Sales:** fell.\n- **Stock:** fine."


def test_generate_digest_with_no_sales_yet_has_no_chart(monkeypatch):
    _install_digest_data(monkeypatch, latest_by_channel={})
    fake_client = _install_fake_client(monkeypatch, [_digest_answer()])

    result = assistant.generate_digest(customer="fairprice")

    assert "no sales data loaded yet" in _facts_sent(fake_client)
    assert result["has_chart"] is False


def test_generate_digest_refusal_still_returns_a_usable_response(monkeypatch):
    _install_digest_data(monkeypatch)
    blocked = FakeResponse(finish_reason="SAFETY", content={"role": "model", "parts": []})
    _install_fake_client(monkeypatch, [blocked])

    result = assistant.generate_digest(customer="fairprice")

    assert result["answer"]
    assert result["has_chart"] is True  # the chart is still built from real data regardless


# ---- catalog outage inside a sales tool --------------------------------------------

def test_sales_tool_database_outage_becomes_a_tool_error_not_a_crash(monkeypatch):
    # Sales tools now read Cloud SQL directly -- an outage there (not just a
    # BigQuery one) must still become a tool error the model is told about,
    # not an unhandled crash, and the turn should still produce an answer.
    from sqlalchemy.exc import OperationalError

    def _outage(**kwargs):
        raise OperationalError("SELECT 1", {}, Exception("connection refused"))

    monkeypatch.setattr(sellout_service, "get_dashboard_summary", _outage)

    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "week"}))
    fake_client = _install_fake_client(
        monkeypatch,
        [tool_call, _final_answer(grounded=True, data_source="none", text="I couldn't load the sales data right now.")],
    )

    result = assistant.ask("how is revenue trending?", history=None, customer="fairprice")

    function_response = _dump(fake_client.models.calls[1]["contents"][-1])["parts"][0]["function_response"]
    assert "sales database" in function_response["response"]["error"]
    assert result["answer"] == "I couldn't load the sales data right now."


# ---- today and the report date ---------------------------------------------------

def test_today_is_the_singapore_day_not_the_utc_day(monkeypatch):
    # 30 Sep 17:00 UTC is 1 Oct 01:00 in Singapore: "this month" is October.
    _freeze_clock(monkeypatch)

    assert assistant._singapore_today() == datetime.date(2026, 10, 1)


def test_system_prompt_gives_the_singapore_date(monkeypatch):
    _freeze_clock(monkeypatch)

    assert "Today's date is 2026-10-01." in assistant._system_prompt(dict(assistant._DEFAULT_CHART_STYLE))


def test_report_date_is_the_singapore_day(monkeypatch):
    _freeze_clock(monkeypatch)

    assert assistant._report_date() == "October 01, 2026"


# ---- ask(): the data shown always belongs to this turn's question -------------
#
# Regression for the supporting table lagging behind the question: a chart
# from an earlier turn was carried over whenever a turn fetched nothing, and
# earlier tool results were replayed to the model, which then answered from
# them. These check the table/chart is only ever this turn's own data.

_INVENTORY_CUSTOMERS = [
    {"customer_id": 1, "customer_name": "fairprice"},
    {"customer_id": 2, "customer_name": "sheng_shiong"},
]


def _customer_view(customer_id, doh, month="2026-07-01", product="Adult Pants L"):
    name = {1: "fairprice", 2: "sheng_shiong"}[customer_id]
    return {
        "customer": {"customer_id": customer_id, "customer_name": name},
        "threshold": {"target_doh": 30.0, "min_doh": 25.0, "max_doh": 35.0, "is_global_default": True},
        "trend": [],
        "skus": [
            {"sku": "S1", "product_name": product, "month": month, "ending_stock": 663,
             "doh": doh, "target_doh": 30.0, "doh_status": "above_max"},
        ],
    }


_PREVIOUS_TABLE = {
    "has_chart": False, "chart_type": "none", "chart_categories": [], "chart_series": [],
    "has_table": True, "table_columns": ["month", "product", "ending_stock", "doh", "target_doh", "status"],
    "table_rows": [["2026-06-01", "Old Product", "100", "10.0", "30.0", "Below min"]],
}

_PROMOTIONS = [{
    "stores": [{"retailer": "fairprice", "store_name": "All"}], "period_label": "June",
    "promo_type": "discount", "promotion_mechanic": "20% off",
}]


def test_an_answer_that_fetched_nothing_does_not_show_the_previous_table(monkeypatch):
    _install_fake_client(monkeypatch, [_final_answer(text="Happy to help with anything else.")])

    result = assistant.ask("thanks", history=None, customer="fairprice", last_chart=_PREVIOUS_TABLE)

    assert result["has_table"] is False
    assert result["table_rows"] == []


def test_a_lookup_that_errors_does_not_fall_back_to_the_previous_table(monkeypatch):
    monkeypatch.setattr(inventory_service, "list_customers", lambda: _INVENTORY_CUSTOMERS)
    tool_call = _tool_call_response(("tu_1", "get_customer_inventory", {"customer": "cold storage"}))
    _install_fake_client(
        monkeypatch,
        [tool_call, _final_answer(text="I could not find a customer called cold storage.")],
    )

    result = assistant.ask("what is cold storage's DOH?", history=None, customer="fairprice", last_chart=_PREVIOUS_TABLE)

    assert result["has_table"] is False


def test_a_follow_up_about_the_chart_on_screen_keeps_it(monkeypatch):
    _install_fake_client(monkeypatch, [_final_answer(text="Because stock built up.", uses_previous_chart=True)])

    result = assistant.ask("why is that?", history=None, customer="fairprice", last_chart=_PREVIOUS_TABLE)

    assert result["table_rows"] == _PREVIOUS_TABLE["table_rows"]


def test_a_fresh_lookup_replaces_the_previous_table_even_if_the_model_calls_it_a_follow_up(monkeypatch):
    monkeypatch.setattr(inventory_service, "list_customers", lambda: _INVENTORY_CUSTOMERS)
    monkeypatch.setattr(inventory_service, "get_customer_view", lambda **kwargs: _customer_view(2, 22.5))
    tool_call = _tool_call_response(("tu_1", "get_customer_inventory", {"customer": "sheng shiong"}))
    _install_fake_client(
        monkeypatch,
        [tool_call, _final_answer(text="Sheng Shiong is at 22.5 days.", uses_previous_chart=True)],
    )

    result = assistant.ask("and sheng shiong?", history=None, customer="fairprice", last_chart=_PREVIOUS_TABLE)

    assert result["table_rows"][0][3] == "22.5"


def test_earlier_tool_calls_and_results_are_not_replayed_to_the_model(monkeypatch):
    fake_client = _install_fake_client(monkeypatch, [_final_answer(text="Okay.")])
    history = [
        {"role": "user", "parts": [{"text": "what is fairprice's DOH?"}]},
        {"role": "model", "parts": [{"function_call": {"id": "tu_1", "name": "get_customer_inventory", "args": {}}}]},
        {"role": "user", "parts": [{"function_response": {"name": "get_customer_inventory", "response": {"result": "{}"}}}]},
        {"role": "model", "parts": [{"text": '{"answer": "39.8 days"}'}]},
    ]

    assistant.ask("and sheng shiong?", history=history, customer="fairprice")

    sent = fake_client.models.calls[0]["contents"]
    assert sent[:-1] == [history[0], history[3]]
    assert not any("function_call" in part or "function_response" in part for c in sent for part in c["parts"])


def test_a_model_turn_keeps_its_text_when_its_tool_call_is_dropped():
    mixed = {"role": "model", "parts": [{"text": "Checking."}, {"function_call": {"id": "tu_1", "name": "x", "args": {}}}]}

    assert assistant._text_only_history([mixed]) == [{"role": "model", "parts": [{"text": "Checking."}]}]


# ---- ask(): more than one call to an inventory tool in one turn ---------------

def test_two_customer_inventory_calls_show_both_customers_rows(monkeypatch):
    monkeypatch.setattr(inventory_service, "list_customers", lambda: _INVENTORY_CUSTOMERS)
    monkeypatch.setattr(
        inventory_service, "get_customer_view",
        lambda customer_id, **kwargs: _customer_view(customer_id, 39.8 if customer_id == 1 else 22.5),
    )
    tool_calls = _tool_call_response(
        ("tu_1", "get_customer_inventory", {"customer": "fairprice"}),
        ("tu_2", "get_customer_inventory", {"customer": "sheng shiong"}),
    )
    _install_fake_client(
        monkeypatch,
        [tool_calls, _final_answer(text="Fairprice is at 39.8 days and Sheng Shiong at 22.5 days.")],
    )

    result = assistant.ask("compare DOH", history=None, customer="fairprice")

    assert result["table_columns"][0] == "customer"
    assert [row[0] for row in result["table_rows"]] == ["fairprice", "sheng_shiong"]
    assert [row[4] for row in result["table_rows"]] == ["39.8", "22.5"]


def test_repeated_customer_inventory_calls_for_one_customer_show_each_row_once(monkeypatch):
    monkeypatch.setattr(inventory_service, "list_customers", lambda: _INVENTORY_CUSTOMERS)
    monkeypatch.setattr(
        inventory_service, "get_customer_view",
        lambda customer_id, **kwargs: _customer_view(customer_id, 39.8, month=kwargs["start_month"].isoformat()),
    )
    tool_calls = _tool_call_response(
        ("tu_1", "get_customer_inventory", {"customer": "fairprice", "start_month": "2026-06-01"}),
        ("tu_2", "get_customer_inventory", {"customer": "fairprice", "start_month": "2026-06-01"}),
        ("tu_3", "get_customer_inventory", {"customer": "fairprice", "start_month": "2026-07-01"}),
    )
    _install_fake_client(monkeypatch, [tool_calls, _final_answer(text="Fairprice is at 39.8 days.")])

    result = assistant.ask("DOH for June and July", history=None, customer="fairprice")

    assert result["table_columns"][0] == "month"
    assert [row[0] for row in result["table_rows"]] == ["2026-06-01", "2026-07-01"]


def test_the_inventory_table_wins_over_promotions_when_both_are_fetched():
    chart = assistant._build_chart(
        [
            ("get_customer_inventory", {}, _customer_view(1, 39.8)),
            ("get_promotions", {}, _PROMOTIONS),
        ],
        "fairprice",
    )

    assert chart["table_columns"] == ["month", "product", "ending_stock", "doh", "target_doh", "status"]
    assert chart["table_rows"][0][3] == "39.8"


def test_promotions_are_shown_when_nothing_else_was_fetched():
    chart = assistant._build_chart([("get_promotions", {}, _PROMOTIONS)], "fairprice")

    assert chart["table_columns"] == ["stores", "period", "type", "mechanic"]


# ---- ask(): the headline figure must be one the table shows -------------------

def _inventory_turn(monkeypatch, doh=39.8):
    monkeypatch.setattr(inventory_service, "list_customers", lambda: _INVENTORY_CUSTOMERS)
    monkeypatch.setattr(inventory_service, "get_customer_view", lambda **kwargs: _customer_view(1, doh))
    return _tool_call_response(("tu_1", "get_customer_inventory", {"customer": "fairprice"}))


def test_a_doh_figure_that_matches_the_table_is_not_questioned(monkeypatch):
    tool_call = _inventory_turn(monkeypatch)
    fake_client = _install_fake_client(monkeypatch, [tool_call, _final_answer(text="Fairprice is at 39.8 days.")])

    result = assistant.ask("what is fairprice's DOH?", history=None, customer="fairprice")

    assert len(fake_client.models.calls) == 2
    assert result["has_table"] is True


def test_a_rounded_doh_figure_still_matches_the_table(monkeypatch):
    tool_call = _inventory_turn(monkeypatch)
    fake_client = _install_fake_client(monkeypatch, [tool_call, _final_answer(text="Fairprice is at about 40 days.")])

    assistant.ask("what is fairprice's DOH?", history=None, customer="fairprice")

    assert len(fake_client.models.calls) == 2


def test_a_doh_figure_that_is_not_in_the_table_triggers_one_correction(monkeypatch):
    tool_call = _inventory_turn(monkeypatch)
    fake_client = _install_fake_client(
        monkeypatch,
        [
            tool_call,
            _final_answer(text="Fairprice is at 31.2 days."),
            _final_answer(text="Fairprice is at 39.8 days."),
        ],
    )

    result = assistant.ask("what is fairprice's DOH?", history=None, customer="fairprice")

    assert result["answer"] == "Fairprice is at 39.8 days."
    assert result["has_table"] is True
    correction = fake_client.models.calls[2]["contents"][-1]["parts"][0]["text"]
    assert "31.2 days" in correction and "39.8" in correction


def test_a_doh_figure_that_stays_wrong_leaves_the_table_out_without_an_error_message(monkeypatch):
    tool_call = _inventory_turn(monkeypatch)
    _install_fake_client(
        monkeypatch,
        [
            tool_call,
            _final_answer(text="Fairprice is at 31.2 days."),
            _final_answer(text="Fairprice is at 33.3 days."),
        ],
    )

    result = assistant.ask("what is fairprice's DOH?", history=None, customer="fairprice")

    assert result["has_table"] is False
    assert result["answer"].startswith("Fairprice is at 31.2 days.")
    assert "verify" not in result["answer"]


def test_a_unit_figure_in_a_sales_answer_is_not_checked_against_inventory(monkeypatch):
    summary = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-01", "revenue": 5473.15, "units": 388}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: summary)
    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"granularity": "month"}))
    fake_client = _install_fake_client(monkeypatch, [tool_call, _final_answer(text="We sold 388 units in August.")])

    result = assistant.ask("units in august", history=None, customer="fairprice")

    assert len(fake_client.models.calls) == 2
    assert result["has_chart"] is True


# ---- ask(): figures quoted from memory are re-fetched --------------------------

def test_quoting_a_figure_without_calling_a_tool_triggers_one_fresh_lookup(monkeypatch):
    tool_call = _inventory_turn(monkeypatch, doh=22.5)
    from_memory = _final_answer(text="Sheng Shiong is at 39.8 days.")
    fresh = _final_answer(text="Fairprice is at 22.5 days.")
    fake_client = _install_fake_client(monkeypatch, [from_memory, tool_call, fresh])

    result = assistant.ask("what about now?", history=None, customer="fairprice")

    assert len(fake_client.models.calls) == 3
    assert "call the relevant tool" in fake_client.models.calls[1]["contents"][-1]["parts"][0]["text"]
    assert result["table_rows"][0][3] == "22.5"
    assert result["answer"] == "Fairprice is at 22.5 days."


def test_a_follow_up_about_the_chart_on_screen_is_not_sent_back_to_fetch(monkeypatch):
    fake_client = _install_fake_client(
        monkeypatch, [_final_answer(text="It is at 10.0 days because stock was run down.", uses_previous_chart=True)]
    )

    assistant.ask("why is that?", history=None, customer="fairprice", last_chart=_PREVIOUS_TABLE)

    assert len(fake_client.models.calls) == 1


def test_a_web_search_answer_shows_no_table_from_earlier_tool_calls(monkeypatch):
    tool_call = _inventory_turn(monkeypatch)
    business_attempt = _final_answer(grounded=False, data_source="none", text="No internal data for that.")
    search_attempt = _search_answer("Competitor news.")
    reshaped = _final_answer(grounded=True, data_source="market_data", text="A competitor launched a product.")
    _install_fake_client(monkeypatch, [tool_call, business_attempt, search_attempt, reshaped])

    result = assistant.ask("any competitor news?", history=None, customer="fairprice")

    assert result["has_table"] is False


def test_a_thank_you_that_quotes_old_figures_gets_a_plain_reply_with_no_table_or_note(monkeypatch):
    from_memory = _final_answer(text="You're welcome! Sheng Shiong is at 81.2 days.")
    plain = _final_answer(text="You're welcome!", grounded=True, data_source="none")
    fake_client = _install_fake_client(monkeypatch, [from_memory, plain])

    result = assistant.ask("thanks", history=None, customer="fairprice", last_chart=_PREVIOUS_TABLE)

    assert "thank-you" in fake_client.models.calls[1]["contents"][-1]["parts"][0]["text"]
    assert result["answer"] == "You're welcome!"
    assert result["has_table"] is False


def test_an_unverifiable_figure_with_no_table_shows_the_answer_unchanged(monkeypatch):
    still_from_memory = _final_answer(text="Sheng Shiong is at 81.2 days.")
    fake_client = _install_fake_client(monkeypatch, [still_from_memory, still_from_memory])

    result = assistant.ask("how are they doing?", history=None, customer="fairprice")

    assert result["answer"] == "Sheng Shiong is at 81.2 days."
    assert result["has_table"] is False


# ---- the inventory table shows the current position, not the oldest rows ------

def _history_view():
    view = _customer_view(1, 5.3, month="2025-01-01", product="Old Month Pants")
    view["skus"] += [
        {"sku": "S1", "product_name": f"Pants {month}", "month": month, "ending_stock": 100,
         "doh": 20.0 + i, "target_doh": 30.0, "doh_status": "below_min"}
        for i, month in enumerate(["2025-02-01", "2026-06-01", "2026-07-01", "2026-07-01"])
    ]
    return view


def test_a_customer_inventory_question_with_no_range_shows_only_the_newest_month():
    chart = assistant._build_chart([("get_customer_inventory", {"customer": "fairprice"}, _history_view())], "fairprice")

    assert {row[0] for row in chart["table_rows"]} == {"2026-07-01"}
    assert len(chart["table_rows"]) == 2


def test_a_customer_inventory_question_with_a_range_shows_the_newest_months_first():
    chart = assistant._build_chart(
        [("get_customer_inventory", {"customer": "fairprice", "start_month": "2025-01-01"}, _history_view())],
        "fairprice",
    )

    months = [row[0] for row in chart["table_rows"]]
    assert months == sorted(months, reverse=True)
    assert months[0] == "2026-07-01"


def test_the_customer_wide_doh_is_in_the_table_above_that_months_skus():
    view = _customer_view(1, 39.8)
    view["trend"] = [{"month": "2026-07-01", "ending_stock": 663, "sell_in": 0, "doh": 33.4,
                      "target_doh": 30.0, "doh_status": "within"}]

    chart = assistant._build_chart([("get_customer_inventory", {"customer": "fairprice"}, view)], "fairprice")

    assert chart["table_rows"][0] == ["2026-07-01", "All products", "663", "33.4", "30.0", "Within range"]
    assert chart["table_rows"][1][1] == "Adult Pants L"


def test_an_answer_quoting_the_customer_wide_doh_keeps_its_table(monkeypatch):
    view = _customer_view(1, 39.8)
    view["trend"] = [{"month": "2026-07-01", "ending_stock": 663, "sell_in": 0, "doh": 33.4,
                      "target_doh": 30.0, "doh_status": "within"}]
    monkeypatch.setattr(inventory_service, "list_customers", lambda: _INVENTORY_CUSTOMERS)
    monkeypatch.setattr(inventory_service, "get_customer_view", lambda **kwargs: view)
    tool_call = _tool_call_response(("tu_1", "get_customer_inventory", {"customer": "fairprice"}))
    fake_client = _install_fake_client(monkeypatch, [tool_call, _final_answer(text="Fairprice's DOH is 33.4 days.")])

    result = assistant.ask("what is fairprice's DOH?", history=None, customer="fairprice")

    assert len(fake_client.models.calls) == 2
    assert result["has_table"] is True


# ---- sales summary: a month with only monthly totals is not "no data" ---------

def test_an_empty_weekly_summary_falls_back_to_monthly(monkeypatch):
    monthly = {
        "offline": {"storeFormats": [], "periodTotal": [{"period_label": "August 2026", "period_start": "2026-08-01", "revenue": 13225.35, "units": 1009}], "periodByFormat": []},
        "online": {"storeFormats": [], "periodTotal": [], "periodByFormat": []},
    }
    empty = {"offline": {"periodTotal": []}, "online": {"periodTotal": []}}
    granularities = []

    def fake_summary(granularity, **kwargs):
        granularities.append(granularity)
        return monthly if granularity == "month" else empty

    monkeypatch.setattr(sellout_service, "get_dashboard_summary", fake_summary)
    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"mode": "offline", "start_date": "2026-08-01", "end_date": "2026-08-31"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="Offline sales in August were $13,225.35.")])

    result = assistant.ask("What were offline sales in August?", history=None, customer="fairprice")

    assert granularities == ["week", "month"]
    assert result["chart_categories"] == ["August 2026"]
    assert result["chart_series"] == [{"label": "Revenue", "values": [13225.35]}]


def test_a_weekly_summary_with_rows_is_not_re_queried_monthly(monkeypatch):
    weekly = {"offline": {"periodTotal": [{"period_label": "Week 1", "period_start": "2026-08-03", "revenue": 100.0, "units": 1}]}, "online": {"periodTotal": []}}
    granularities = []

    def fake_summary(granularity, **kwargs):
        granularities.append(granularity)
        return weekly

    monkeypatch.setattr(sellout_service, "get_dashboard_summary", fake_summary)
    tool_call = _tool_call_response(("tu_1", "get_sales_summary", {"mode": "offline"}))
    _install_fake_client(monkeypatch, [tool_call, _final_answer(text="Week 1 was $100.00.")])

    assistant.ask("weekly sales", history=None, customer="fairprice")

    assert granularities == ["week"]


# ---- a restyle keeps the original answer for the PDF/PowerPoint ----------------

_REVENUE_CHART = {
    "has_chart": True, "chart_type": "bar", "chart_categories": ["August 2026"],
    "chart_series": [{"label": "Revenue", "values": [13225.35]}],
    "has_table": False, "table_columns": [], "table_rows": [],
}
_ORIGINAL_ANSWER = "Offline sales in August 2026 were $13,225.35 across 1,009 units."


def test_a_restyle_exports_the_original_answer_not_the_confirmation(monkeypatch):
    red = {**assistant._DEFAULT_CHART_STYLE, "color": "#DC2626"}
    _install_fake_client(
        monkeypatch, [_final_answer(text="The chart is now red.", uses_previous_chart=True, chart_style=red)]
    )

    result = assistant.ask(
        "make it red", history=None, customer="fairprice",
        last_chart=_REVENUE_CHART, last_answer=_ORIGINAL_ANSWER,
    )

    assert result["answer"] == "The chart is now red."
    assert result["export_answer"] == _ORIGINAL_ANSWER
    assert result["chart_series"] == _REVENUE_CHART["chart_series"]


def test_a_follow_up_explanation_exports_its_own_answer(monkeypatch):
    _install_fake_client(
        monkeypatch, [_final_answer(text="Because a promotion ran.", uses_previous_chart=True)]
    )

    result = assistant.ask(
        "why is that?", history=None, customer="fairprice",
        last_chart=_REVENUE_CHART, last_answer=_ORIGINAL_ANSWER,
    )

    assert result["export_answer"] == "Because a promotion ran."


def test_a_new_question_exports_its_own_answer(monkeypatch):
    _install_fake_client(monkeypatch, [_final_answer(text="Hello!", data_source="none")])

    result = assistant.ask("hi", history=None, customer="fairprice", last_answer=_ORIGINAL_ANSWER)

    assert result["export_answer"] == "Hello!"


# ---- customer names the way people write them, and restyles labelled "none" ---

def test_a_customer_written_with_spaces_and_capitals_resolves_to_its_stored_name(monkeypatch):
    monkeypatch.setattr(
        inventory_service, "list_customers",
        lambda: [{"customer_id": 1, "customer_name": "fairprice"}, {"customer_id": 9, "customer_name": "sheng_shiong"}],
    )

    assert assistant._resolve_inventory_customer_id("Sheng Shiong") == 9
    assert assistant._resolve_inventory_customer_id("sheng-shiong") == 9
    assert assistant._resolve_inventory_customer_id("sheng_shiong") == 9


def test_a_sales_tool_receives_the_customer_as_a_slug(monkeypatch):
    seen = {}
    monkeypatch.setattr(sellout_service, "get_dashboard_summary", lambda **kwargs: seen.update(kwargs) or {})

    assistant._run_business_tool("get_sales_summary", {"customer": "Sheng Shiong"}, "fairprice")

    assert seen["customer"] == "sheng_shiong"


def test_a_restyle_keeps_the_chart_even_when_the_model_labels_it_data_source_none(monkeypatch):
    red = {**assistant._DEFAULT_CHART_STYLE, "color": "#DC2626"}
    _install_fake_client(
        monkeypatch,
        [_final_answer(text="Updated the colour.", data_source="none", uses_previous_chart=True, chart_style=red)],
    )

    result = assistant.ask(
        "make it red", history=None, customer="fairprice",
        last_chart=_REVENUE_CHART, last_answer=_ORIGINAL_ANSWER,
    )

    assert result["chart_series"] == _REVENUE_CHART["chart_series"]
    assert result["export_answer"] == _ORIGINAL_ANSWER


def test_the_doh_table_is_shown_when_the_stock_overview_was_fetched_too():
    overview = {"skus": [{"month": "2026-08-01", "customer_name": "fairprice", "product_name": "Pants", "ending_stock": 500}]}
    view = _customer_view(1, 39.8)

    chart = assistant._build_chart(
        [("get_inventory_overview", {}, overview), ("get_customer_inventory", {"customer": "fairprice"}, view)],
        "fairprice",
    )

    assert "doh" in chart["table_columns"]


def test_the_stock_overview_is_still_shown_when_it_is_the_only_inventory_tool_used():
    overview = {"skus": [{"month": "2026-08-01", "customer_name": "fairprice", "product_name": "Pants", "ending_stock": 500}]}

    chart = assistant._build_chart([("get_inventory_overview", {}, overview)], "fairprice")

    assert chart["table_columns"] == ["month", "customer", "product", "ending_stock"]


def test_a_thank_you_wrongly_flagged_as_about_the_chart_does_not_show_it(monkeypatch):
    _install_fake_client(
        monkeypatch,
        [_final_answer(text="You're welcome!", data_source="none", uses_previous_chart=True)],
    )

    result = assistant.ask(
        "thanks", history=None, customer="fairprice",
        last_chart=_REVENUE_CHART, last_answer=_ORIGINAL_ANSWER,
    )

    assert result["has_chart"] is False
    assert result["export_answer"] == "You're welcome!"


# ---- a range total asked across several customers ------------------------------

def _range_total(customer, revenue):
    return (
        {"customer": customer, "current_start": "2026-08-01", "current_end": "2026-08-31"},
        {"current": {"start": "2026-08-01", "end": "2026-08-31", "revenue": revenue, "units": 1},
         "previous": {"start": "2026-07-01", "end": "2026-07-31", "revenue": 1.0, "units": 1}},
    )


def test_a_range_total_for_several_customers_shows_one_bar_per_customer():
    sources = [("compare_periods", *_range_total(c, r)) for c, r in
               (("fairprice", 24052.28), ("nhg", 28155.5), ("sheng_siong", 0.0))]

    chart = assistant._build_chart(sources, "fairprice")

    assert chart["chart_categories"] == ["Fairprice", "Nhg", "Sheng Siong"]
    assert chart["chart_series"] == [{"label": "Revenue", "values": [24052.28, 28155.5, 0.0]}]


def test_a_comparison_across_several_customers_shows_the_pages_customer():
    fair_input, fair_raw = _range_total("fairprice", 24052.28)
    nhg_input, nhg_raw = _range_total("nhg", 28155.5)
    fair_input = {**fair_input, "previous_start": "2026-07-01", "previous_end": "2026-07-31"}
    nhg_input = {**nhg_input, "previous_start": "2026-07-01", "previous_end": "2026-07-31"}

    chart = assistant._build_chart(
        [("compare_periods", fair_input, fair_raw), ("compare_periods", nhg_input, nhg_raw)], "fairprice"
    )

    assert chart["chart_categories"] == ["Previous", "Current"]
    assert chart["chart_series"][0]["values"] == [1.0, 24052.28]
