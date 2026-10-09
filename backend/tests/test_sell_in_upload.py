from datetime import date
from io import BytesIO

import openpyxl
import pytest

from app.services.sell_in_upload import product_for_description, quantity_problem, read_sell_in_lines

HEADER = [
    "Month of Sales",
    "Invoice Year 2026",
    "Customer Name",
    "STORE",
    "PO NO.",
    "DESCRIPTION",
    "QTY in Packs",
]


def _workbook(rows, header=HEADER, sheet="Data Entry SI"):
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = sheet
    ws.append(header)
    for row in rows:
        ws.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_reads_the_month_of_sales_in_the_invoice_year_with_its_quantity_in_packs():
    data = _workbook([["Aug", 2026, "Fairprice", "Joo Koon", "PO1", "AIRE ADULT PANTS S/M 10S", 168]])

    lines = read_sell_in_lines(data)

    assert lines == [
        {
            "customer_name": "Fairprice",
            "month": date(2026, 8, 1),
            "description": "AIRE ADULT PANTS S/M 10S",
            "qty": 168.0,
        }
    ]


@pytest.mark.parametrize("spelling", ["Sep", "Sept"])
def test_september_is_read_in_either_spelling(spelling):
    data = _workbook([[spelling, 2026, "Fairprice", "x", "PO1", "AIRE ADULT PANTS L 10S", 8]])

    assert read_sell_in_lines(data)[0]["month"] == date(2026, 9, 1)


def test_an_unreadable_month_or_quantity_is_read_as_none_not_guessed():
    data = _workbook(
        [
            ["Octobr", 2026, "Fairprice", "x", "PO1", "AIRE ADULT PANTS L 10S", 8],
            ["Aug", 2026, "Fairprice", "x", "PO2", "AIRE ADULT PANTS L 10S", None],
        ]
    )

    lines = read_sell_in_lines(data)

    assert lines[0]["month"] is None
    assert lines[1]["qty"] is None


def test_blank_rows_are_left_out():
    data = _workbook([[None, None, None, None, None, None, None], ["Aug", 2026, "Fairprice", "x", "PO1", "AIRE ADULT PANTS L 10S", 8]])

    assert len(read_sell_in_lines(data)) == 1


def test_the_sheet_is_found_by_its_columns_whatever_it_is_called():
    data = _workbook([["Aug", 2026, "Fairprice", "x", "PO1", "AIRE ADULT PANTS L 10S", 8]], sheet="Sheet7")

    assert read_sell_in_lines(data)[0]["customer_name"] == "Fairprice"


def test_the_header_may_sit_below_a_title_block():
    book = openpyxl.Workbook()
    ws = book.active
    ws.append(["Sell-in tracker", None])
    ws.append([None, "Updated Sep 2026"])
    ws.append(HEADER)
    ws.append(["Aug", 2026, "Fairprice", "x", "PO1", "AIRE ADULT PANTS L 10S", 8])
    buffer = BytesIO()
    book.save(buffer)

    lines = read_sell_in_lines(buffer.getvalue())

    assert lines[0]["qty"] == 8.0


def test_a_file_with_no_sheet_holding_the_columns_is_refused():
    data = _workbook([], header=[h for h in HEADER if h != "QTY in Packs"])

    with pytest.raises(ValueError, match="No sheet in this file has the sell-in columns"):
        read_sell_in_lines(data)


def test_two_sheets_with_the_columns_are_refused_rather_than_guessed():
    book = openpyxl.Workbook()
    book.active.title = "Jan"
    book.active.append(HEADER)
    second = book.create_sheet("Feb")
    second.append(HEADER)
    buffer = BytesIO()
    book.save(buffer)

    with pytest.raises(ValueError, match="More than one sheet"):
        read_sell_in_lines(buffer.getvalue())


@pytest.mark.parametrize(
    "description, product",
    [
        ("AIRE ADULT PANTS S/M 10S", "Aire Adult Pants S/M"),
        ("AIRE ADULT PANTS XL 10S", "Aire Adult Pants XL"),
        ("Adult Pull Up Pants Ultra L", "Aire Ultra Pants L"),
        ("Adult Pull Up Pants Ultra XL", "Aire Ultra Pants XL"),
        ("Adult Tape Ultra S/M", "Aire Ultra Tape S/M"),
    ],
)
def test_each_tracker_description_maps_to_its_catalog_product(description, product):
    assert product_for_description(description) == product


def test_a_description_that_is_not_listed_is_not_guessed():
    assert product_for_description("AIRE ADULT PANTS XL 10S - Ultra") is None
    assert product_for_description("") is None


def test_a_sheet_with_only_a_header_is_refused_with_a_message_about_rows():
    data = _workbook([])

    with pytest.raises(ValueError, match="has no rows to upload"):
        read_sell_in_lines(data)


def test_a_completely_empty_workbook_is_refused_as_having_no_sell_in_columns():
    book = openpyxl.Workbook()
    book.active.title = "Data Entry SI"
    buffer = BytesIO()
    book.save(buffer)

    with pytest.raises(ValueError, match="No sheet in this file has the sell-in columns"):
        read_sell_in_lines(buffer.getvalue())


@pytest.mark.parametrize(
    "qty, problem",
    [
        (None, "Quantity in packs is missing"),
        (-8.0, "Quantity in packs is negative"),
        (0.0, None),
        (168.0, None),
    ],
)
def test_a_quantity_must_be_present_and_not_negative(qty, problem):
    assert quantity_problem(qty) == problem


def test_a_file_that_is_not_an_excel_workbook_is_refused_with_a_plain_message():
    csv_text = "Month of Sales,Invoice Year 2026,Customer Name" + '\\n' + "Aug,2026,Fairprice" + '\\n'
    with pytest.raises(ValueError, match="^Invalid file. Please upload the correct file.$"):
        read_sell_in_lines(csv_text.encode())
