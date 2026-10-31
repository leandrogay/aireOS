from openpyxl import Workbook

from scripts.import_reviewed_promotions import (
    _candidate_key,
    _existing_keys,
    _load_review,
    _payload,
)
from scripts.preview_promotion_calendars import preview_calendar


def test_preview_reads_filled_monthly_and_weekly_offers_without_template_rows(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheng Siong Online"
    sheet.append(["MONTHLY REGULAR PROMOTIONS"])
    sheet.append(["period_start", "period_end", "sku_range", "promotion_mechanic",
                  "period_label", "voucher", "promotion_mechanic_2", "carton_mechanic"])
    sheet.append(["2026-01-01", "2026-01-31", "Flagship", "20% Off",
                  "Jan-2026", "$3 off $30", "Buy 2 Get 30% Off", "30% Off"])
    sheet.append(["2026-02-01", "2026-02-28", "Ultra Pants", None,
                  "Feb-2026", None, None])
    sheet.append(["week_start", "week_end", "sku_range", "side_offer_mechanic",
                  "week_label", "voucher"])
    sheet.append(["2026-01-08", "2026-01-14", "Flagship", "Buy 1 Get 1 Free",
                  "W2-2026", None])
    # Excel normally stores dates as datetimes, not date strings.
    from datetime import datetime
    for row in (3, 4, 6):
        sheet.cell(row, 1).value = datetime.fromisoformat(sheet.cell(row, 1).value)
        sheet.cell(row, 2).value = datetime.fromisoformat(sheet.cell(row, 2).value)
    path = tmp_path / "calendar.xlsx"
    workbook.save(path)

    rows = preview_calendar(path)
    assert len(rows) == 2
    assert {item["promo_type"] for item in rows} == {
        "monthly", "side_offer",
    }
    assert rows[0]["mechanic_column"] == "promotion_mechanic_2"
    assert all(item["retailer_proposal"] == "sheng_siong_online" for item in rows)
    assert all(item["approved"] is False for item in rows)
    assert rows[0]["estimated_multiplier"] == "0.765"


def test_review_csv_requires_explicit_approval(tmp_path):
    path = tmp_path / "review.csv"
    path.write_text(
        "approved,issues,voucher,period_label,sku_range_proposal\n"
        "FALSE,,$3 off $30,Jan-2026,Range A\n"
        "TRUE,,,,Range B\n",
        encoding="utf-8",
    )
    rows = _load_review(path)
    assert [row["approved"] for row in rows] == [False, True]
    assert rows[1]["voucher"] is None
    assert rows[1]["issues"] == []


def test_singhealth_customer_scope_links_both_channels_once(tmp_path):
    from datetime import datetime

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SingHealth"
    sheet.append(["period_start", "period_end", "sku_range", "promotion_mechanic"])
    sheet.append([datetime(2026, 9, 1), datetime(2026, 9, 30), "Flagship", "20% Off"])
    path = tmp_path / "singhealth.xlsx"
    workbook.save(path)

    row = preview_calendar(path)[0]
    assert row["retailer_proposal"] == "singhealth"
    stores = [
        {"retailer_name": "singhealth_offline", "store_code": "81001",
         "store_name": "BEDOK", "store_format": None},
        {"retailer_name": "singhealth_online", "store_code": "88001",
         "store_name": "ONLINE", "store_format": None},
    ]
    payload = _payload(row, stores)
    assert {store.retailer for store in payload.stores} == {
        "singhealth_offline", "singhealth_online",
    }
    existing = [{
        "period_start": row["period_start"],
        "period_end": row["period_end"],
        "promo_type": row["promo_type"],
        "promotion_mechanic": row["promotion_mechanic"],
        "voucher": row["voucher"],
        "stores": [{"retailer": store["retailer_name"], "store_code": store["store_code"]}
                   for store in stores],
        "skus": [{"sku_range": row["sku_range_proposal"]}],
    }]
    assert _candidate_key(row, stores) in _existing_keys(existing)
    nhg_store = [{"retailer_name": "nhg_offline", "store_code": "NHG",
                  "store_name": "NHG", "store_format": None}]
    assert _candidate_key(row, nhg_store)[:-1] != _candidate_key(row, stores)[:-1]
