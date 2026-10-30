"""
Golden tests of mapping generation on the real retailer sample files
(mapping hand-off, sections 3 and 8).

The files are commercial data and are never committed: put them in
tests/fixtures/retailer_samples/ (git-ignored). Without them every test here
is skipped -- CI covers the same logic on the synthetic copies in conftest.

Each file's reference contract is given the layout detection finds, validated
against the file's headers and checked statically against the file and the
hand-off's SKU master. The last section then runs every sample file through
the transformation and reconciles what it would load against the totals the
file states for itself.
"""

import io
import re
from pathlib import Path

import openpyxl
import pytest

from app.services import (
    column_profile,
    layout_detection,
    mapping_checks,
    mapping_contract,
    reference_contracts,
    sellout_ingestion,
    validation_service,
)

SAMPLES = Path(__file__).parent / "fixtures" / "retailer_samples"
# Searched for anywhere under SAMPLES: the files sit in one folder per retailer.
PATTERNS = {
    "fairprice_monthly": "XEL_VENDORS_BRANDS_MONTH_*",
    "fairprice_weekly": "XEL_VENDORS_BRANDS_WEEK_*",
    "nhg_sap": "Aire*Sell*Out*",
    "sheng_siong_claim": "Claim_Sales_Summary_Report_*",
    "singhealth_er130": "XEL [A-Z][a-z][a-z]2[0-9]*.xlsx",
    "prime_jkdc": "XEL Lifecare JKDC Withdrawal Report*",
}
# The hand-off's SKU master (section 7).
MASTER = {
    "13255043", "13255044", "13255045", "13271337", "13271338",
    "13271339", "13271340", "13271341", "13271342",
}


def _paths(name: str) -> list[Path]:
    return sorted(SAMPLES.rglob(PATTERNS[name])) if SAMPLES.exists() else []


# Every sample file of every layout, for the tests that run them all.
ALL_SAMPLES = [(name, path) for name in PATTERNS for path in _paths(name)]


def _sample(name: str, latest: bool = False) -> tuple[str, bytes]:
    """The first sample file of a layout, or its latest export (the FairPrice tests count its months)."""
    matches = _paths(name)
    if not matches:
        pytest.skip(f"sample file for {name} not in {SAMPLES}")
    path = matches[-1] if latest else matches[0]
    return path.name, path.read_bytes()


def _contract_for(name: str, filename: str, data: bytes):
    """The reference contract made for this file, as generation would make it."""
    frame, layout = layout_detection.read_table(filename, data)
    proposed = reference_contracts.reference_contract(name)
    proposed["layout"] = mapping_contract.layout_instructions(layout)
    contract = mapping_contract.validate_contract(proposed, list(frame.columns), trust_review=True)
    checks = mapping_checks.run_checks(frame, layout, contract, filename, MASTER)
    return frame, layout, contract, {check["id"]: check for check in checks}


def _workbook(data: bytes):
    return openpyxl.load_workbook(io.BytesIO(data))


def _saved(workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---- Every sample ---------------------------------------------------------------------

@pytest.mark.parametrize("name", list(PATTERNS))
def test_the_reference_contract_validates_cleanly_and_passes_every_check(name):
    _, _, contract, checks = _contract_for(name, *_sample(name))

    assert contract["warnings"] == []
    assert mapping_checks.failures(list(checks.values())) == []
    assert contract["aggregate"] == reference_contracts.reference_contract(name).get("aggregate")


@pytest.mark.parametrize("name, codes", [
    ("fairprice_monthly", 9), ("nhg_sap", 9), ("sheng_siong_claim", 6), ("singhealth_er130", 9), ("prime_jkdc", 8),
])
def test_every_product_code_reads_as_a_master_sku(name, codes):
    frame, _, contract, checks = _contract_for(name, *_sample(name))

    found = mapping_checks.sku_codes(frame, contract)
    assert len(found) == codes
    assert {item["sku"] for item in found} <= MASTER
    assert checks["sku_codes"]["passed"] is True


# ---- 3.1 / 8.1 FairPrice ------------------------------------------------------------------

def test_fairprice_layout_and_raw_grain():
    frame, layout, contract, checks = _contract_for("fairprice_monthly", *_sample("fairprice_monthly", latest=True))

    assert (layout["header_row"], len(frame)) == (1, 166)
    assert not frame.duplicated(["SKU No.", "Store Code", "Sales UOM"]).any()
    assert contract["layout"]["header_row"] == 1
    sales, qty = (mapping_contract.resolve_selector(g["select"], list(frame.columns)) for g in contract["melt_groups"])
    assert len(sales) == len(qty) == 8
    assert checks["period"]["passed"] is True
    assert checks["rules_retailer"]["evidence"]["matched"] == ["FPON"]


# ---- 3.2 / 8.2 NHG -----------------------------------------------------------------------

def test_nhg_reads_the_data_sheet_not_the_pivot():
    frame, layout, contract, _ = _contract_for("nhg_sap", *_sample("nhg_sap"))

    assert layout["sheet"] == "Sheet1"
    assert layout["rejected_sheets"][0]["reason"] == "pivot table"
    assert len(frame) == 291
    assert contract["layout"]["sheet"] == "Sheet1"
    assert contract["layout"]["sheet_position"] == 2
    # Kept for the transformation to reconcile against: -1 x this is 1,868.
    assert layout["pivot_totals"][0]["total"] == -1868


def test_nhg_quantities_are_negated():
    _, _, contract, checks = _contract_for("nhg_sap", *_sample("nhg_sap"))

    assert contract["fields"]["quantity_units"]["transforms"] == [{"type": "negate"}]
    assert checks["quantity_sign"]["passed"] is True
    assert contract["fields"]["store_code"]["status"] == "not_available"
    assert contract["fields"]["revenue"]["status"] == "not_available"


# ---- 3.3 / 8.3 Sheng Siong -----------------------------------------------------------------

def test_sheng_siong_layout():
    frame, layout, contract, _ = _contract_for("sheng_siong_claim", *_sample("sheng_siong_claim"))

    assert layout["header_row"] == 7
    assert layout["metadata"]["RefNo"] == "P100045355"
    assert layout["metadata"]["Supplier Code"] == "XLPL"
    assert (layout["stated_period"]["start"], layout["stated_period"]["end"]) == ("2026-06-01", "2026-06-30")
    assert len(frame) == 123
    assert [row["label"] for row in layout["footer_rows"]] == ["Total", "Total Claim before GST", "Total Claim After GST"]
    assert contract["layout"]["header_row"] == 7
    assert contract["layout"]["fill_down"] == ["BusinessDate"]


def test_sheng_siong_melts_the_20_outlets_and_never_reads_the_claim_as_revenue():
    frame, _, contract, checks = _contract_for("sheng_siong_claim", *_sample("sheng_siong_claim"))

    outlets = mapping_contract.resolve_selector(contract["melt_groups"][0]["select"], list(frame.columns))
    assert len(outlets) == 20 and "OUTLET" not in outlets and "Total" not in outlets
    assert checks["rules_retailer"]["evidence"]["matched"] == ["6M_Ecomm", "KL_Ecomm"]
    assert contract["fields"]["revenue"]["status"] == "not_available"
    assert contract["data_source"] == "promo_claim"


# ---- 8.4 Robustness variants -----------------------------------------------------------------

def test_a_month_shifted_fairprice_file_has_the_same_layout_and_contract():
    filename, data = _sample("fairprice_monthly", latest=True)

    def later(match):
        day, month, year = match.groups()
        return f"{day}-{int(month) + 1:02d}-{year}"

    header, _, body = data.decode("utf-8-sig").partition("\n")
    shifted = (re.sub(r"(\d{2})-(\d{2})-(\d{4})", later, header) + "\n" + body).encode("utf-8")
    shifted_name = "XEL_VENDORS_BRANDS_MONTH_01-02-2026_01-09-2026_1.txt"

    original, _, _, _ = _contract_for("fairprice_monthly", filename, data)
    frame, _, contract, checks = _contract_for("fairprice_monthly", shifted_name, shifted)

    assert column_profile.fingerprint(list(original.columns)) == column_profile.fingerprint(list(frame.columns))
    sales = mapping_contract.resolve_selector(contract["melt_groups"][0]["select"], list(frame.columns))
    assert len(sales) == 8 and sales[-1].endswith("01-09-2026")
    assert mapping_checks.failures(list(checks.values())) == []


def test_an_extra_sheng_siong_outlet_is_selected_by_the_same_contract():
    filename, data = _sample("sheng_siong_claim")
    workbook = _workbook(data)
    sheet = workbook.active
    total_column = next(cell.column for cell in sheet[7] if str(cell.value).strip() == "Total")
    sheet.insert_cols(total_column)
    sheet.cell(7, total_column, "ZZ1")
    variant = _saved(workbook)

    original, _, _, _ = _contract_for("sheng_siong_claim", filename, data)
    frame, _, contract, checks = _contract_for("sheng_siong_claim", filename, variant)

    assert column_profile.fingerprint(list(original.columns)) == column_profile.fingerprint(list(frame.columns))
    outlets = mapping_contract.resolve_selector(contract["melt_groups"][0]["select"], list(frame.columns))
    assert len(outlets) == 21 and "ZZ1" in outlets
    assert mapping_checks.failures(list(checks.values())) == []


def test_nhg_with_its_sheets_swapped_still_reads_the_data_sheet():
    filename, data = _sample("nhg_sap")
    workbook = _workbook(data)
    workbook.move_sheet("Sheet1", offset=-1)

    frame, layout, contract, _ = _contract_for("nhg_sap", filename, _saved(workbook))

    assert layout["sheet"] == "Sheet1"
    assert len(frame) == 291
    assert contract["layout"]["sheet_position"] == 1


def test_an_unseen_nhg_code_fails_the_product_code_check():
    filename, data = _sample("nhg_sap")
    workbook = _workbook(data)
    sheet = workbook["Sheet1"]
    sheet.append([sheet.cell(2, 1).value, "0059-01-099-R", "AIRE ULTRA PROTECT ADULT TAPE XL 12S", -5])

    _, _, _, checks = _contract_for("nhg_sap", filename, _saved(workbook))

    skus = checks["sku_codes"]
    assert skus["passed"] is False
    assert [item["code"] for item in skus["evidence"]["unknown"]] == ["0059-01-099-R"]


# ---- Transforming every sample ----------------------------------------------------------

def _transformed(name: str, path: Path):
    """What loading this file through its reference contract would store, and its layout."""
    data = path.read_bytes()
    frame, layout, contract, _ = _contract_for(name, path.name, data)
    rows = sellout_ingestion.transform_confirmed_upload(path.name, data, contract)
    return validation_service.validate_mapped_dataframe(rows), frame, layout


def _of(name: str) -> list[Path]:
    return [path for key, path in ALL_SAMPLES if key == name]


@pytest.mark.parametrize("name, path", ALL_SAMPLES, ids=[path.name for _, path in ALL_SAMPLES])
def test_every_sample_file_transforms_with_no_row_rejected(name, path):
    validated, _, _ = _transformed(name, path)

    assert validated["rows_ingested"] > 0
    assert validated["rejection_summary"] == ""


@pytest.mark.parametrize("path", _of("sheng_siong_claim"), ids=lambda path: path.name)
def test_sheng_siong_quantities_equal_the_statements_total_row(path):
    validated, _, layout = _transformed("sheng_siong_claim", path)

    total = next(row for row in layout["footer_rows"] if row["label"] == "Total")
    assert validated["valid_df"]["quantity_units"].sum() == float(total["values"]["Total"])


@pytest.mark.parametrize("path", _of("nhg_sap"), ids=lambda path: path.name)
def test_nhg_quantities_equal_minus_the_pivot_grand_total(path):
    validated, _, layout = _transformed("nhg_sap", path)

    assert validated["valid_df"]["quantity_units"].sum() == -layout["pivot_totals"][0]["total"]


@pytest.mark.parametrize("path", _of("singhealth_er130"), ids=lambda path: path.name)
def test_singhealth_units_and_value_equal_the_outlet_grid(path):
    validated, frame, layout = _transformed("singhealth_er130", path)
    rows = validated["valid_df"]

    outlets = mapping_contract.resolve_selector({"kind": "between", "after": "Promo cost"}, list(frame.columns))
    units = frame[outlets].astype(float)
    price = frame[reference_contracts.SINGHEALTH_PRICE_COLUMN].astype(float)
    assert rows["quantity_units"].sum() == units.to_numpy().sum()
    assert rows["revenue"].sum() == pytest.approx(units.mul(price, axis=0).to_numpy().sum())
    assert set(rows["period_start"]) == {layout["stated_period"]["start"]}
    codes = next(item["cells"] for item in layout["rows_above_header"] if item["row"] == layout["header_row"] - 1)
    assert dict(zip(rows["store_name"], rows["store_code"])) == {name: codes[name] for name in set(rows["store_name"])}


@pytest.mark.parametrize("path", _of("prime_jkdc"), ids=lambda path: path.name)
def test_prime_quantities_equal_the_packs_shipped(path):
    validated, frame, _ = _transformed("prime_jkdc", path)

    assert validated["valid_df"]["quantity_units"].sum() == frame["Total Shipped Qty"].astype(float).sum()
